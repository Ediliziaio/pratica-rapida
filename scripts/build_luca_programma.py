from pathlib import Path
from collections import Counter
import csv
import math

from PIL import Image, ImageChops, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
ASSETS = OUT / "assets_luca"
OUT.mkdir(exist_ok=True)
ASSETS.mkdir(exist_ok=True)

DOCX_PATH = OUT / "Luca_Viviani_Proposta_Coordinamento_MB.docx"
LOGO_SOURCE = Path("/private/tmp/futuro_nazionale_logo.png")
EU_LISTE = Path("/private/tmp/europee_liste.csv")
EU_INSIEME = Path("/private/tmp/europee_insieme.csv")

NAVY = "061B3A"
BLUE = "2357A6"
BLUE_LIGHT = "E8F0FC"
RED = "C83E4D"
RED_LIGHT = "FBEAEC"
GREEN = "168C3A"
GOLD = "F0C419"
INK = "172033"
MUTED = "667085"
LIGHT = "F3F5F8"
MID = "D9DEE7"
WHITE = "FFFFFF"
GRAY = "7B8494"
TEAL = "317C7D"

TOWNS = [
    ("Arcore", "CDX", 2021, 50.86, "oltre 15mila"),
    ("Biassono", "CDX", 2021, 43.32, "fino a 15mila"),
    ("Bovisio-Masciago", "COMM", 2024, 50.60, "oltre 15mila"),
    ("Briosco", "CDX", 2021, 54.46, "fino a 15mila"),
    ("Carnate", "CDX", 2022, 39.38, "fino a 15mila"),
    ("Cesano Maderno", "CSX", 2022, 53.94, "oltre 15mila"),
    ("Lentate sul Seveso", "CDX", 2022, 67.03, "oltre 15mila"),
    ("Lesmo", "CSX", 2022, 33.38, "fino a 15mila"),
    ("Limbiate", "CDX", 2021, 71.27, "oltre 15mila"),
    ("Lissone", "CDX", 2022, 50.43, "oltre 15mila"),
    ("Meda", "CDX", 2022, 68.50, "oltre 15mila"),
    ("Monza", "CSX", 2022, 51.21, "oltre 15mila"),
    ("Seveso", "CDX", 2021, 53.52, "oltre 15mila"),
    ("Sulbiate", "CIV", 2022, 53.82, "fino a 15mila"),
    ("Varedo", "CDX", 2021, 68.83, "fino a 15mila"),
    ("Vedano al Lambro", "CDX", 2021, 55.88, "fino a 15mila"),
    ("Verano Brianza", "CSX", 2021, 49.56, "fino a 15mila"),
    ("Vimercate", "CSX", 2021, 60.87, "oltre 15mila"),
]

NAME_MAP = {name.upper(): name for name, *_ in TOWNS}


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_run(run, size=10.2, color=INK, bold=False, italic=False, font="Aptos"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic
    return run


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def borders(cell, color=MID, size=6, sides=("top", "left", "bottom", "right")):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for side in sides:
        edge = tc_borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            tc_borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(size))
        edge.set(qn("w:color"), color)


def set_table_fixed(table, widths_cm):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    total = int(sum(widths_cm) / 2.54 * 1440)
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_cm:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width / 2.54 * 1440)))
        grid.append(col)
    for row in table.rows:
        for idx, (cell, width) in enumerate(zip(row.cells, widths_cm)):
            cell.width = Cm(width)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(int(width / 2.54 * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_p(container, text="", size=10.2, color=INK, bold=False, italic=False,
          align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=5, line=1.15,
          keep=False):
    p = container.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep
    set_run(p.add_run(text), size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1, kicker=None):
    if kicker:
        add_p(doc, kicker.upper(), size=8.4, color=RED, bold=True, after=2, keep=True)
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), size={1: 20, 2: 14, 3: 11.5}[level],
            color=NAVY if level < 3 else BLUE, bold=True)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    set_run(p.add_run(text), size=9.7)
    return p


STEP_NO = 0


def reset_steps():
    global STEP_NO
    STEP_NO = 0


def add_number(doc, text):
    global STEP_NO
    STEP_NO += 1
    table = doc.add_table(rows=1, cols=2)
    set_table_fixed(table, [0.85, 16.95])
    shade(table.cell(0, 0), RED)
    shade(table.cell(0, 1), WHITE)
    borders(table.cell(0, 0), WHITE, 6)
    borders(table.cell(0, 1), WHITE, 6)
    cell_text(table.cell(0, 0), str(STEP_NO), size=8.8, color=WHITE, bold=True,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(table.cell(0, 1), text, size=9.6)
    return table


def cell_text(cell, text, size=9.2, color=INK, bold=False,
              align=WD_ALIGN_PARAGRAPH.LEFT, after=0):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    set_run(p.add_run(text), size=size, color=color, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margins(cell)
    return p


def add_page_break(doc):
    doc.add_page_break()


def add_metric_strip(doc, items):
    table = doc.add_table(rows=1, cols=len(items))
    set_table_fixed(table, [17.8 / len(items)] * len(items))
    for cell, (value, label) in zip(table.rows[0].cells, items):
        shade(cell, NAVY)
        borders(cell, NAVY)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        set_run(p.add_run(value), size=17, color=GOLD, bold=True)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        set_run(p2.add_run(label.upper()), size=7.6, color=WHITE, bold=True)
        cell_margins(cell, top=130, bottom=130)
    return table


def add_callout(doc, title, body, fill=BLUE_LIGHT, accent=BLUE):
    table = doc.add_table(rows=1, cols=2)
    set_table_fixed(table, [0.16, 17.64])
    shade(table.cell(0, 0), accent)
    shade(table.cell(0, 1), fill)
    borders(table.cell(0, 0), accent)
    borders(table.cell(0, 1), fill)
    cell_text(table.cell(0, 0), "")
    c = table.cell(0, 1)
    c.text = ""
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run(title), size=10.2, color=NAVY, bold=True)
    p2 = c.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.12
    set_run(p2.add_run(body), size=9.3, color=INK)
    cell_margins(c, top=120, bottom=120, start=150, end=150)
    return table


def crop_logo():
    img = Image.open(LOGO_SOURCE).convert("RGBA")
    threshold = img.convert("RGB").point(lambda p: 255 if p < 242 else 0).convert("L")
    bbox = threshold.getbbox()
    if bbox:
        pad = 18
        bbox = (max(0, bbox[0]-pad), max(0, bbox[1]-pad),
                min(img.width, bbox[2]+pad), min(img.height, bbox[3]+pad))
        img = img.crop(bbox)
    mask = Image.new("L", img.size, 0)
    md = ImageDraw.Draw(mask)
    inset = max(4, int(min(img.size) * .01))
    md.ellipse((inset, inset, img.width - inset, img.height - inset), fill=255)
    img.putalpha(mask)
    out = ASSETS / "logo_cropped.png"
    img.save(out)
    return out


def read_eu_data():
    wanted = set(NAME_MAP)
    code_to_town = {}
    turnout = {}
    with EU_INSIEME.open(encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            name = row["DESCRIZIONE COMUNE"]
            if row["SIGLA"] == "MB" and name in wanted:
                code_to_town[row["CODICE ISTAT"]] = NAME_MAP[name]
                turnout[NAME_MAP[name]] = float(row["perc_vot"].replace(",", "."))
    party = {town: {} for town in code_to_town.values()}
    votes = Counter()
    with EU_LISTE.open(encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            town = code_to_town.get(row["CODICE ISTAT"])
            if not town:
                continue
            pct = float(row["perc"].replace(",", "."))
            party[town][row["desc_lis"]] = pct
            votes[row["desc_lis"]] += int(row["voti"])
    out = {}
    for town, p in party.items():
        fdi = p.get("FRATELLI D'ITALIA", 0)
        lega = p.get("LEGA SALVINI PREMIER", 0)
        fi = p.get("FORZA ITALIA - NOI MODERATI - PPE", 0)
        pd = p.get("PARTITO DEMOCRATICO", 0)
        avs = p.get("ALLEANZA VERDI E SINISTRA", 0)
        m5s = p.get("MOVIMENTO 5 STELLE", 0)
        out[town] = {
            "turnout": turnout[town], "cdx": fdi + lega + fi,
            "csx": pd + avs, "m5s": m5s,
            "fdi": fdi, "pd": pd, "lega": lega, "fi": fi, "avs": avs,
        }
    total = sum(votes.values())
    aggregate = {
        "FdI": 100 * votes["FRATELLI D'ITALIA"] / total,
        "PD": 100 * votes["PARTITO DEMOCRATICO"] / total,
        "FI–NM": 100 * votes["FORZA ITALIA - NOI MODERATI - PPE"] / total,
        "Lega": 100 * votes["LEGA SALVINI PREMIER"] / total,
        "AVS": 100 * votes["ALLEANZA VERDI E SINISTRA"] / total,
        "M5S": 100 * votes["MOVIMENTO 5 STELLE"] / total,
    }
    return out, aggregate


def chart_parties(aggregate):
    labels = list(aggregate)
    values = [aggregate[k] for k in labels]
    colors = ["173F8A", "C83E4D", "2B72B8", "3C8BCC", "D96A73", "7D6BAF"]
    img = Image.new("RGB", (1500, 680), "white")
    d = ImageDraw.Draw(img)
    bold = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 38)
    regular = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 28)
    small_bold = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 28)
    d.text((50, 35), "Europee 2024 · voto aggregato nei 18 comuni al voto", font=bold, fill="#061B3A")
    maxv = 34.0
    for i, (label, value, color) in enumerate(zip(labels, values, colors)):
        y = 125 + i * 84
        d.text((50, y + 10), label, font=regular, fill="#172033")
        x0, x1 = 300, 300 + int(920 * value / maxv)
        d.rounded_rectangle((x0, y, x1, y + 54), radius=12, fill="#" + color)
        d.text((x1 + 18, y + 9), f"{value:.1f}%".replace('.', ','), font=small_bold, fill="#172033")
    path = ASSETS / "europee_partiti.png"
    img.save(path)
    return path


def chart_admin():
    values = [12, 5, 1]
    colors = ["#2357A6", "#C83E4D", "#317C7D"]
    img = Image.new("RGB", (820, 680), "white")
    d = ImageDraw.Draw(img)
    bold = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 34)
    mid = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 24)
    big = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 62)
    d.text((40, 30), "Ultima maggioranza comunale vincente", font=bold, fill="#061B3A")
    box = (190, 125, 630, 565)
    start = -90
    for value, color in zip(values, colors):
        end = start + 360 * value / sum(values)
        d.arc(box, start=start, end=end, fill=color, width=90)
        start = end
    d.text((350, 285), "18", font=big, fill="#061B3A", anchor="mm")
    d.text((350, 340), "COMUNI", font=mid, fill="#667085", anchor="mm")
    x = 110
    for label, color in zip(["12 CDX*", "5 CSX", "1 CIVICA"], colors):
        d.rounded_rectangle((x, 610, x+28, 638), radius=5, fill=color)
        d.text((x+38, 609), label, font=mid, fill="#172033")
        x += 220
    path = ASSETS / "amministrative_18.png"
    img.save(path)
    return path


def page_cover(doc, logo):
    table = doc.add_table(rows=1, cols=2)
    set_table_fixed(table, [12.7, 5.1])
    for c in table.rows[0].cells:
        shade(c, NAVY)
        borders(c, NAVY)
        cell_margins(c, top=180, bottom=180, start=180, end=180)
    left, right = table.rows[0].cells
    left.text = ""
    p = left.paragraphs[0]
    set_run(p.add_run("PROPOSTA DI COORDINAMENTO PROVINCIALE"), size=9.2, color=GOLD, bold=True)
    p2 = left.add_paragraph()
    p2.paragraph_format.space_before = Pt(4)
    set_run(p2.add_run("MONZA E BRIANZA"), size=12, color=WHITE, bold=True)
    right.text = ""
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run().add_picture(str(logo), width=Cm(2.7))

    add_p(doc, "BRIANZA", size=37, color=NAVY, bold=True, before=80, after=0)
    add_p(doc, "2027", size=52, color=RED, bold=True, after=12)
    add_p(doc, "Organizzare. Radicare. Preparare.", size=18, color=NAVY, bold=True, after=18)
    add_p(doc, "Proposta politica e organizzativa di Luca Viviani", size=14, color=INK, bold=True, after=5)
    add_p(doc, "per il coordinamento provinciale di Futuro Nazionale", size=11.5, color=MUTED, after=38)
    add_metric_strip(doc, [("55", "comuni MB"), ("18", "al voto 2027"),
                           ("≈435mila", "cittadini coinvolti"), ("9", "mesi alla sfida")])
    add_p(doc, "BOZZA RISERVATA · CONFRONTO INTERNO", size=8.2, color=MUTED, bold=True,
          before=42, after=2)
    add_p(doc, "28 agosto 2026 · Dati aggiornati alle fonti disponibili", size=8.2, color=MUTED)


def page_executive(doc):
    add_heading(doc, "Una guida pronta a trasformare crescita in struttura", kicker="Sintesi esecutiva")
    add_p(doc, "Futuro Nazionale è nella fase decisiva del radicamento. Monza e Brianza non parte da zero: esistono comitati, aderenti, amministratori ed energie che chiedono un coordinamento riconoscibile, ordinato e capace di preparare le amministrative della primavera 2027.", size=11.2, line=1.25, after=12)
    add_callout(doc, "La proposta", "Un coordinamento provinciale che unisca le realtà esistenti, assegni responsabilità verificabili, dia voce ai territori e consegni entro 100 giorni una macchina organizzativa pronta a lavorare sui 55 comuni, con priorità immediata sui 18 chiamati al voto.")
    add_heading(doc, "Perché Luca Viviani", level=2)
    strengths = [
        ("Esperienza pubblica", "Ha ricoperto ruoli di consigliere provinciale e assessore comunale: conosce istituzioni, amministratori e tempi reali degli enti locali."),
        ("Competenza e formazione", "Architetto e docente per una vita negli istituti superiori: unisce lettura del territorio, capacità didattica e metodo."),
        ("Tempo operativo", "Oggi pensionato, può garantire presenza continuativa durante il giorno, non soltanto nelle ore serali."),
        ("Leadership di squadra", "Propone deleghe vere, obiettivi scritti e verifiche periodiche: inclusione significa responsabilità, non moltiplicazione di titoli."),
    ]
    table = doc.add_table(rows=2, cols=2)
    set_table_fixed(table, [8.9, 8.9])
    for cell, (title, body) in zip([c for r in table.rows for c in r.cells], strengths):
        shade(cell, LIGHT)
        borders(cell, WHITE, 10)
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        set_run(p.add_run(title), size=10.2, color=NAVY, bold=True)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        set_run(p2.add_run(body), size=9.2, color=INK)
        cell_margins(cell, top=140, bottom=140, start=150, end=150)
    add_heading(doc, "Quattro impegni misurabili", level=2)
    reset_steps()
    for text in [
        "Mappa completa di persone, comitati, amministratori e presìdi entro 30 giorni dalla nomina.",
        "Assetto provinciale e cinque aree operative attive entro 60 giorni, in coerenza con lo Statuto.",
        "Dossier politico-organizzativo per ciascuno dei 18 comuni al voto entro 90 giorni.",
        "Rendiconto mensile sintetico: attività svolte, tesseramento, copertura territoriale, criticità e decisioni richieste.",
    ]:
        add_number(doc, text)


def page_alignment(doc):
    add_heading(doc, "Coerenza statutaria, identità chiara, disciplina operativa", kicker="Mandato")
    add_callout(doc, "La carica corretta", "Lo Statuto usa la denominazione Coordinatore Provinciale. Nella fase transitoria, gli organi di coordinamento provinciale sono nominati dal Coordinatore Regionale d’intesa con il Coordinatore Nazionale; successivamente il Coordinatore è eletto dal Congresso Provinciale e dura in carica tre anni.", fill="FFF7D9", accent=GOLD)
    add_heading(doc, "V.I.T.A.L.E. tradotto in comportamenti", level=2)
    rows = [
        ("V · Virtù", "Parola data, puntualità, rendicontazione e responsabilità personale."),
        ("I · Identità", "Una linea politica riconoscibile, collegata alla storia e alle comunità brianzole."),
        ("T · Tradizioni", "Presenza nei paesi, nelle associazioni e nei luoghi reali della vita locale."),
        ("A · Amore", "Cura della comunità, delle famiglie e delle fragilità senza assistenzialismo."),
        ("L · Libertà", "Delega operativa dentro regole chiare; iniziativa rapida sulle decisioni reversibili."),
        ("E · Eccellenza", "Incarichi per competenza, obiettivi misurabili e sostituzione di ciò che non funziona."),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_fixed(table, [3.3, 14.5])
    for i, (lab, desc) in enumerate(rows):
        shade(table.cell(i, 0), NAVY)
        shade(table.cell(i, 1), LIGHT if i % 2 == 0 else WHITE)
        borders(table.cell(i, 0), WHITE)
        borders(table.cell(i, 1), MID)
        cell_text(table.cell(i, 0), lab, size=9.2, color=WHITE, bold=True)
        cell_text(table.cell(i, 1), desc, size=9.2)
    add_heading(doc, "Principi di conduzione", level=2)
    for text in [
        "Una sola linea politica provinciale; molte responsabilità operative distribuite.",
        "Nessun comitato cancellato o scavalcato: ascolto, riconoscimento del lavoro svolto, raccordo e standard comuni.",
        "Le nomine statutarie seguono lo Statuto; le deleghe operative non creano organi paralleli.",
        "Accordi elettorali, candidature, comunicazioni ad alto impatto e impegni economici restano al livello competente e richiedono tracciabilità.",
    ]:
        add_bullet(doc, text)


def page_org(doc):
    add_heading(doc, "Una squadra larga, con responsabilità nette", kicker="Architettura organizzativa")
    add_p(doc, "Il modello distingue gli organi previsti dallo Statuto dalle deleghe operative. Ogni funzione ha un responsabile unico, un risultato atteso, indicatori e una verifica mensile.", after=10)
    top = doc.add_table(rows=1, cols=1)
    set_table_fixed(top, [17.8])
    shade(top.cell(0, 0), NAVY)
    borders(top.cell(0, 0), NAVY)
    cell_text(top.cell(0, 0), "COORDINATORE PROVINCIALE · indirizzo politico, rappresentanza, direzione", size=11, color=WHITE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p(doc, "", after=1)
    table = doc.add_table(rows=2, cols=3)
    set_table_fixed(table, [5.93, 5.94, 5.93])
    statutory = [
        ("Esecutivo Provinciale", "10 componenti quando costituito secondo Statuto; riunione almeno mensile."),
        ("Organizzazione", "Eventi, sedi, calendario, logistica, copertura territoriale."),
        ("Tesseramento", "Registro, crescita, rinnovi e integrità del dato."),
        ("Direzione disciplinare", "Garanzie, regole e gestione delle controversie."),
        ("Futuro Nazionale Giovani", "Scuola, università e nuova classe dirigente."),
        ("Segreteria di coordinamento", "Agenda, verbali, cruscotto e follow-up; funzione tecnica non statutaria."),
    ]
    for cell, (title, body) in zip([c for r in table.rows for c in r.cells], statutory):
        shade(cell, BLUE_LIGHT)
        borders(cell, WHITE, 10)
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        set_run(p.add_run(title), size=9.7, color=NAVY, bold=True)
        p2 = cell.add_paragraph()
        set_run(p2.add_run(body), size=8.7)
        cell_margins(cell, top=130, bottom=130, start=135, end=135)
    add_heading(doc, "Deleghe operative proponibili", level=2)
    roles = [
        ("5 referenti d’area", "copertura dei comuni e raccordo con i comitati"),
        ("Enti locali", "relazione con amministratori e temi sovracomunali"),
        ("Amministrative 2027", "dossier, scadenze e coordinamento tecnico"),
        ("Programma e consulte", "laboratori tematici e proposte territoriali"),
        ("Comunicazione", "messaggi, contenuti e calendario editoriale"),
        ("Eventi e piazza", "permessi, materiali, presìdi e sicurezza"),
        ("Dati e monitoraggio", "analisi aggregate, cruscotto e report"),
        ("Compliance e privacy", "regole, consensi, archivi e conflitti d’interesse"),
    ]
    table = doc.add_table(rows=4, cols=2)
    set_table_fixed(table, [8.9, 8.9])
    for cell, (title, body) in zip([c for r in table.rows for c in r.cells], roles):
        shade(cell, LIGHT)
        borders(cell, WHITE, 8)
        cell.text = ""
        p = cell.paragraphs[0]
        set_run(p.add_run(title + " · "), size=9.2, color=BLUE, bold=True)
        set_run(p.add_run(body), size=9.0)
        cell_margins(cell, top=105, bottom=105)
    add_p(doc, "Nota: per ora il documento indica soltanto funzioni. Nomi, perimetri e deleghe saranno definiti dopo l’ascolto e con l’assenso degli organi competenti.", size=8.3, color=MUTED, italic=True, before=5)


def page_areas(doc):
    add_heading(doc, "Cinque aree operative, una sola provincia", kicker="Presenza territoriale")
    add_p(doc, "Le aree sono strumenti di lavoro, non nuovi livelli statutari. Servono a ridurre le distanze, distribuire il carico e garantire una presenza regolare. Il perimetro definitivo sarà validato dopo il censimento dei 55 comuni.", after=10)
    areas = [
        ("01 · Monza e cintura", "Monza · Vedano al Lambro", "Capoluogo, servizi, mobilità, sicurezza urbana"),
        ("02 · Ovest e asse del Seveso", "Bovisio-Masciago · Cesano Maderno · Lentate · Limbiate · Meda · Seveso · Varedo", "Sette comuni al voto, infrastrutture e tessuto produttivo"),
        ("03 · Brianza centrale", "Biassono · Lissone", "Reti civiche, commercio, associazionismo e prossimità"),
        ("04 · Valle del Lambro", "Briosco · Verano Brianza", "Comunità locali, territorio e collegamenti"),
        ("05 · Vimercatese", "Arcore · Carnate · Lesmo · Sulbiate · Vimercate", "Innovazione, imprese, mobilità e presidio dell’Est"),
    ]
    table = doc.add_table(rows=len(areas), cols=3)
    set_table_fixed(table, [4.4, 7.5, 5.9])
    for i, (name, towns, focus) in enumerate(areas):
        shade(table.cell(i, 0), NAVY)
        shade(table.cell(i, 1), LIGHT if i % 2 == 0 else WHITE)
        shade(table.cell(i, 2), LIGHT if i % 2 == 0 else WHITE)
        for j in range(3): borders(table.cell(i, j), MID)
        cell_text(table.cell(i, 0), name, size=9.0, color=WHITE, bold=True)
        cell_text(table.cell(i, 1), towns, size=8.8)
        cell_text(table.cell(i, 2), focus, size=8.8, color=MUTED)
    add_heading(doc, "Standard minimo di ogni area", level=2)
    standards = [
        ("Agenda", "una riunione operativa mensile e un calendario condiviso"),
        ("Copertura", "un contatto verificato per ogni comune; priorità ai 18 al voto"),
        ("Ascolto", "un incontro territoriale mensile con verbale e tre problemi prioritari"),
        ("Comitati", "accompagnamento al riconoscimento quando ci sono almeno 10 iscritti"),
        ("Rendiconto", "report di una pagina: fatti, numeri, rischi, decisioni richieste"),
    ]
    table = doc.add_table(rows=1, cols=5)
    set_table_fixed(table, [3.56] * 5)
    for cell, (title, body) in zip(table.rows[0].cells, standards):
        shade(cell, BLUE_LIGHT)
        borders(cell, WHITE, 8)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(title.upper()), size=8.2, color=BLUE, bold=True)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p2.add_run(body), size=8.2)
        cell_margins(cell, top=120, bottom=120)


def page_timeline(doc):
    add_heading(doc, "Dal mandato alla macchina pronta", kicker="Cronoprogramma")
    add_p(doc, "T0 indica la data di nomina. Se T0 coincidesse con il 1° settembre 2026, i primi 100 giorni terminerebbero il 10 dicembre 2026.", size=9.2, color=MUTED, italic=True, after=10)
    stages = [
        ("T0–15", "Ascoltare", "Incontri con comitati e referenti; raccolta contatti, attività, criticità e disponibilità."),
        ("T0–30", "Censire", "Mappa dei 55 comuni; fotografia di iscritti, amministratori, presìdi e competenze."),
        ("T0–60", "Organizzare", "Funzioni statutarie, cinque aree operative, calendario provinciale e regole di lavoro."),
        ("T0–90", "Preparare", "Diciotto dossier comunali; priorità politiche; fabbisogni organizzativi e scadenze."),
        ("Gen–Feb 2027", "Selezionare", "Proposte locali, interlocuzioni e decisioni elettorali ai livelli competenti."),
        ("Mar–Mag 2027", "Eseguire", "Cabina elettorale provinciale, supporto ai comuni e verifica settimanale."),
    ]
    table = doc.add_table(rows=len(stages), cols=3)
    set_table_fixed(table, [3.2, 3.3, 11.3])
    for i, (when, verb, body) in enumerate(stages):
        shade(table.cell(i, 0), RED if i < 4 else NAVY)
        shade(table.cell(i, 1), BLUE_LIGHT)
        shade(table.cell(i, 2), LIGHT if i % 2 == 0 else WHITE)
        for j in range(3): borders(table.cell(i, j), WHITE, 8)
        cell_text(table.cell(i, 0), when, size=9.0, color=WHITE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(table.cell(i, 1), verb, size=9.2, color=NAVY, bold=True)
        cell_text(table.cell(i, 2), body, size=8.9)
    add_heading(doc, "Cruscotto mensile", level=2)
    metrics = [
        ("Copertura", "% comuni con referente verificato"),
        ("Radicamento", "iscritti, rinnovi e comitati riconosciuti"),
        ("Attività", "riunioni, eventi e presenza nei territori"),
        ("Preparazione", "dossier comunali completi / 18"),
        ("Disciplina", "azioni scadute e decisioni ferme"),
        ("Qualità", "partecipazione, follow-up e rispetto delle regole"),
    ]
    table = doc.add_table(rows=2, cols=3)
    set_table_fixed(table, [5.93, 5.94, 5.93])
    for cell, (title, body) in zip([c for r in table.rows for c in r.cells], metrics):
        shade(cell, LIGHT)
        borders(cell, WHITE, 8)
        cell.text = ""
        p = cell.paragraphs[0]
        set_run(p.add_run(title), size=9.2, color=BLUE, bold=True)
        p2 = cell.add_paragraph()
        set_run(p2.add_run(body), size=8.6, color=MUTED)


def page_political(doc):
    add_heading(doc, "Otto laboratori per una proposta brianzola", kicker="Agenda politica territoriale")
    add_p(doc, "Il coordinamento non deve sostituire il programma nazionale, ma tradurlo in proposte locali verificabili. Ogni laboratorio produce entro febbraio 2027 una scheda breve con problema, dati, competenze istituzionali, soluzione e costo indicativo.", after=10)
    labs = [
        ("Sicurezza e legalità", "presìdi, degrado, commercio, polizia locale integrata"),
        ("Impresa e lavoro", "PMI, artigianato, burocrazia, formazione e attrattività"),
        ("Famiglia e comunità", "servizi, casa, anziani, natalità, volontariato e fragilità"),
        ("Scuola e giovani", "merito, orientamento, ITS, sport, cittadinanza e competenze"),
        ("Territorio e ambiente", "urbanistica, consumo di suolo, bonifiche, acqua e parchi"),
        ("Mobilità e infrastrutture", "ferro, strade, trasporto locale e collegamenti est–ovest"),
        ("Identità e cultura", "storia locale, tradizioni, patrimonio, biblioteche ed eventi"),
        ("Salute e sociale", "medicina territoriale, disabilità, tempi e integrazione dei servizi"),
    ]
    table = doc.add_table(rows=4, cols=2)
    set_table_fixed(table, [8.9, 8.9])
    for idx, (cell, (title, body)) in enumerate(zip([c for r in table.rows for c in r.cells], labs)):
        shade(cell, BLUE_LIGHT if idx % 2 == 0 else LIGHT)
        borders(cell, WHITE, 10)
        cell.text = ""
        p = cell.paragraphs[0]
        set_run(p.add_run(title), size=9.8, color=NAVY, bold=True)
        p2 = cell.add_paragraph()
        set_run(p2.add_run(body), size=8.9)
        cell_margins(cell, top=130, bottom=130)
    add_heading(doc, "Metodo di produzione delle proposte", level=2)
    reset_steps()
    for text in [
        "Partire da un problema osservabile e dalla competenza effettiva di Comune, Provincia o Regione.",
        "Consultare amministratori, professionisti, associazioni e cittadini con pluralità di fonti.",
        "Distinguere il principio politico dalla soluzione tecnica e dai relativi vincoli.",
        "Pubblicare solo proposte approvate, con fonti e responsabilità chiare.",
    ]:
        add_number(doc, text)
    add_callout(doc, "Il vantaggio", "Le consulte non diventano un elenco di incarichi: diventano una fabbrica di proposte, talenti e nuova classe dirigente.")


def page_governance(doc):
    add_heading(doc, "Delegare senza perdere direzione", kicker="Regole di governo")
    add_p(doc, "La squadra funziona quando sa chi decide, entro quali limiti e con quale controllo. Il modello seguente protegge insieme velocità, unità politica e reputazione.", after=10)
    rows = [
        ("Decisioni operative e reversibili", "Delegate al responsabile di funzione entro mandato, calendario e risorse approvate.", "Report mensile"),
        ("Comunicazioni pubbliche provinciali", "Linea condivisa; approvazione del Coordinatore e raccordo con i livelli superiori.", "Archivio messaggi"),
        ("Accordi, candidature e alleanze", "Nessuna promessa autonoma; decisione degli organi competenti secondo Statuto.", "Verbale e nulla osta"),
        ("Spese e raccolta fondi", "Solo procedure autorizzate, tracciabili e conformi alle regole nazionali.", "Rendiconto"),
        ("Dati personali", "Minimizzazione, consenso, accessi per ruolo e separazione fra dati associativi e istituzionali.", "Registro accessi"),
        ("Incarichi", "Competenze, obiettivi, durata e verifica; dichiarazione dei conflitti d’interesse.", "Revisione 90 giorni"),
    ]
    table = doc.add_table(rows=1, cols=3)
    set_table_fixed(table, [4.2, 9.4, 4.2])
    hdr = table.rows[0].cells
    for c, t in zip(hdr, ["TIPO DI DECISIONE", "REGOLA", "PROVA"]):
        shade(c, NAVY); borders(c, WHITE); cell_text(c, t, size=8.3, color=WHITE, bold=True)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, text in enumerate(row):
            shade(cells[j], LIGHT if i % 2 == 0 else WHITE)
            borders(cells[j], MID)
            cell_text(cells[j], text, size=8.55, bold=(j == 0), color=NAVY if j == 0 else INK)
    add_heading(doc, "Una linea etica non negoziabile", level=2)
    add_callout(doc, "Nessun sistema di compensazioni", "Le responsabilità interne servono a produrre risultati politici e organizzativi. Non si promettono candidature, nomine o posti in enti e società. Ogni futura proposta dovrà poggiare su competenza, trasparenza, assenza di conflitti e decisione dell’organo competente.", fill=RED_LIGHT, accent=RED)


def page_scenario(doc, admin_chart, party_chart):
    add_heading(doc, "Il campo di gioco: 18 comuni, una sfida provinciale", kicker="Scenario elettorale 2027")
    add_metric_strip(doc, [("10", "comuni sopra 15mila"), ("8", "comuni sotto 15mila"),
                           ("12–5–1", "CDX · CSX · civica*"), ("50,8%", "area CDX · Europee")])
    table = doc.add_table(rows=1, cols=2)
    set_table_fixed(table, [8.9, 8.9])
    for c in table.rows[0].cells:
        borders(c, WHITE)
        cell_margins(c, top=100, bottom=50)
    table.cell(0, 0).paragraphs[0].add_run().add_picture(str(admin_chart), width=Cm(7.4))
    table.cell(0, 1).paragraphs[0].add_run().add_picture(str(party_chart), width=Cm(8.6))
    add_p(doc, "* Bovisio-Masciago è oggi commissariata dopo le dimissioni del sindaco: nel conteggio 12–5–1 è classificata secondo l’ultima maggioranza vincente di centrodestra. Stato attuale: 11 amministrazioni CDX, 5 CSX, 1 civica e 1 commissario.", size=8.0, color=MUTED, italic=True, after=5)
    add_p(doc, "Metodo Europee 2024: area CDX = FdI + Lega + FI–Noi Moderati; area CSX = PD + AVS. Sono somme descrittive, non coalizioni formalmente presenti sulla scheda. Futuro Nazionale non era una lista concorrente: il dato è un benchmark territoriale, non una stima del suo consenso.", size=8.0, color=MUTED, italic=True, after=10)
    add_heading(doc, "Tre evidenze operative", level=2)
    for text in [
        "La maggioranza dei comuni presenta un bacino europeo favorevole all’area di centrodestra, ma le amministrative premiano candidati, coalizioni e radicamento locale.",
        "Monza, Cesano Maderno, Lesmo e Verano Brianza sono amministrate dal centrosinistra pur mostrando nel 2024 una prevalenza aggregata dell’area CDX: organizzazione e offerta locale fanno la differenza.",
        "Vimercate è l’unico dei 18 comuni in cui il benchmark PD+AVS supera FdI+Lega+FI; richiede un progetto locale particolarmente credibile e autonomo dagli automatismi nazionali.",
    ]:
        add_bullet(doc, text)


def municipality_card(cell, town, status, year, muni_pct, size, eu):
    color = {"CDX": BLUE, "CSX": RED, "CIV": TEAL, "COMM": GRAY}[status]
    light = {"CDX": BLUE_LIGHT, "CSX": RED_LIGHT, "CIV": "E7F3F2", "COMM": "ECEEF2"}[status]
    labels = {"CDX": "CENTRODESTRA", "CSX": "CENTROSINISTRA", "CIV": "CIVICA", "COMM": "COMMISSARIO · ultima vittoria CDX"}
    shade(cell, WHITE); borders(cell, color, 12); cell_margins(cell, top=0, bottom=120, start=80, end=80)
    cell.text = ""
    default_p = cell.paragraphs[0]._element
    cell._tc.remove(default_p)
    head = cell.add_table(rows=1, cols=1)
    set_table_fixed(head, [7.95])
    head.alignment = WD_TABLE_ALIGNMENT.CENTER
    shade(head.cell(0, 0), color); borders(head.cell(0, 0), color)
    head.cell(0, 0).text = ""
    p = head.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(town.upper()), size=9.6, color=WHITE, bold=True)
    p2 = head.cell(0, 0).add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    set_run(p2.add_run(labels[status] + " · " + size), size=7.1, color=WHITE, bold=True)
    cell_margins(head.cell(0, 0), top=100, bottom=100, start=120, end=120)
    p = cell.add_paragraph()
    p.paragraph_format.left_indent = Cm(.12)
    p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run(f"Comunali {year}"), size=8.2, color=MUTED, bold=True)
    set_run(p.add_run(f"   {muni_pct:.2f}%".replace('.', ',')), size=11.8, color=color, bold=True)
    p = cell.add_paragraph()
    p.paragraph_format.left_indent = Cm(.12)
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run("Europee 2024"), size=8.2, color=MUTED, bold=True)
    t = cell.add_table(rows=2, cols=3)
    set_table_fixed(t, [2.65, 2.65, 2.65])
    for r in t.rows:
        for c in r.cells:
            shade(c, light); borders(c, WHITE, 6); cell_margins(c, top=65, bottom=65, start=60, end=60)
    for c, label in zip(t.rows[0].cells, ["AREA CDX", "AREA CSX", "AFFLUENZA"]):
        cell_text(c, label, size=6.8, color=MUTED, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for c, val, col in zip(t.rows[1].cells, [eu['cdx'], eu['csx'], eu['turnout']], [BLUE, RED, INK]):
        cell_text(c, f"{val:.1f}%".replace('.', ','), size=10.2, color=col, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    p = cell.add_paragraph()
    p.paragraph_format.left_indent = Cm(.12)
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(0)
    top = "FdI" if eu['fdi'] >= eu['pd'] else "PD"
    topv = max(eu['fdi'], eu['pd'])
    set_run(p.add_run(f"Primo tra FdI/PD: {top} {topv:.1f}%".replace('.', ',')), size=7.6, color=MUTED)


def pages_municipalities(doc, eu_data):
    for page_index in range(3):
        add_heading(doc, "I 18 comuni al voto", kicker=f"Schede territoriali · {page_index + 1}/3")
        add_p(doc, "Colore della scheda = ultima maggioranza comunale vincente; percentuali europee come benchmark aggregato.", size=8.4, color=MUTED, italic=True, after=7)
        subset = TOWNS[page_index*6:(page_index+1)*6]
        table = doc.add_table(rows=3, cols=3)
        set_table_fixed(table, [8.45, .9, 8.45])
        for r_idx, row in enumerate(table.rows):
            gap = row.cells[1]
            shade(gap, WHITE); borders(gap, WHITE); cell_margins(gap, top=0, bottom=0, start=0, end=0)
            for cell, town in zip((row.cells[0], row.cells[2]), subset[r_idx*2:r_idx*2+2]):
                municipality_card(cell, *town, eu_data[town[0]])
        if page_index < 2:
            add_page_break(doc)


def page_strategy(doc, eu_data):
    add_heading(doc, "Dai numeri alle priorità organizzative", kicker="Lettura strategica")
    strong = sorted([t for t in TOWNS if eu_data[t[0]]["cdx"] >= 55], key=lambda x: -eu_data[x[0]]["cdx"])
    competitive = sorted([t for t in TOWNS if 45 <= eu_data[t[0]]["cdx"] < 55], key=lambda x: -eu_data[x[0]]["cdx"])
    complex_ = sorted([t for t in TOWNS if eu_data[t[0]]["cdx"] < 45], key=lambda x: -eu_data[x[0]]["cdx"])
    groups = [
        ("BACINO CDX FORTE", strong, BLUE, "Radicamento, selezione della classe dirigente e capacità di stare dentro la competizione locale."),
        ("AREA COMPETITIVA", competitive, TEAL, "Presenza continuativa, alleanze locali e proposta concreta possono spostare l’esito."),
        ("CONTESTO COMPLESSO", complex_, RED, "Serve lavoro anticipato, reputazione locale e un’offerta distinta dagli automatismi nazionali."),
    ]
    table = doc.add_table(rows=3, cols=3)
    set_table_fixed(table, [4.1, 7.3, 6.4])
    for i, (label, towns, color, implication) in enumerate(groups):
        shade(table.cell(i, 0), color)
        shade(table.cell(i, 1), LIGHT if i % 2 == 0 else WHITE)
        shade(table.cell(i, 2), LIGHT if i % 2 == 0 else WHITE)
        for j in range(3): borders(table.cell(i, j), WHITE, 8)
        cell_text(table.cell(i, 0), label, size=9.0, color=WHITE, bold=True)
        names = " · ".join(f"{t[0]} {eu_data[t[0]]['cdx']:.1f}%".replace('.', ',') for t in towns)
        cell_text(table.cell(i, 1), names, size=8.5)
        cell_text(table.cell(i, 2), implication, size=8.5, color=MUTED)
    add_p(doc, "Le tre fasce sono una lettura organizzativa basata sul solo benchmark europeo, non una previsione elettorale. Le dinamiche comunali possono divergere in modo significativo.", size=8.2, color=MUTED, italic=True, before=4)
    add_heading(doc, "Priorità fino a dicembre 2026", level=2)
    priorities = [
        ("1", "Unificare", "riconoscere i presìdi esistenti, chiudere sovrapposizioni e definire un calendario comune"),
        ("2", "Coprire", "assegnare un referente operativo a ciascuno dei 18 comuni al voto"),
        ("3", "Conoscere", "produrre dossier con storia elettorale, problemi, reti civiche e fabbisogni"),
        ("4", "Formare", "preparare portavoce, organizzatori, rappresentanti di lista e futuri amministratori"),
        ("5", "Decidere", "portare ai livelli competenti opzioni chiare, rischi e raccomandazioni"),
    ]
    table = doc.add_table(rows=1, cols=5)
    set_table_fixed(table, [3.56] * 5)
    for cell, (num, verb, body) in zip(table.rows[0].cells, priorities):
        shade(cell, NAVY); borders(cell, WHITE, 8); cell.text = ""
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(num), size=18, color=GOLD, bold=True)
        p2 = cell.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p2.add_run(verb.upper()), size=8.5, color=WHITE, bold=True)
        p3 = cell.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p3.add_run(body), size=7.7, color=WHITE)
        cell_margins(cell, top=120, bottom=120)


def page_pitch(doc):
    add_heading(doc, "La proposta da portare al tavolo", kicker="Chiusura")
    add_p(doc, "Luca Viviani non propone una candidatura fondata su un titolo, ma un mandato di costruzione con scadenze, responsabilità e verifica.", size=12, color=NAVY, bold=True, line=1.25, after=12)
    add_heading(doc, "Il mandato richiesto", level=2)
    reset_steps()
    for text in [
        "Riconoscimento del ruolo di Coordinatore Provinciale secondo le procedure transitorie previste dallo Statuto.",
        "Novanta giorni per completare censimento, assetto operativo e dossier dei 18 comuni, con verifica a 30 e 90 giorni.",
        "Facoltà di proporre deleghe operative e referenti d’area, senza anticipare nomi e nel rispetto degli organi statutari.",
        "Accesso ai dati associativi strettamente necessari, alle linee organizzative nazionali e ai contatti ufficiali, nel rispetto della privacy.",
    ]:
        add_number(doc, text)
    add_heading(doc, "Intervento breve", level=2)
    add_callout(doc, "«La Brianza non ha bisogno di un coordinatore che occupi uno spazio, ma di una guida che costruisca una squadra. Porto esperienza amministrativa e provinciale, competenza tecnica, conoscenza del territorio e soprattutto la disponibilità di tempo necessaria. Nei primi cento giorni consegnerò una struttura leggibile, cinque aree operative e diciotto dossier per i comuni al voto. Le responsabilità saranno distribuite, gli obiettivi scritti, i risultati verificati. Chiedo un mandato per unire ciò che già esiste e trasformarlo in una presenza provinciale forte, disciplinata e pronta per il 2027.»", "Versione da 60–75 secondi, volutamente priva di riferimenti a nomi o accordi personali.", fill="FFF7D9", accent=GOLD)
    add_heading(doc, "Le prime tre mosse dopo la nomina", level=2)
    for title, body in [
        ("Entro 48 ore", "messaggio unitario a comitati e referenti: ascolto, continuità e calendario dei primi incontri"),
        ("Entro 7 giorni", "riunione operativa provinciale con mappa preliminare e raccolta delle disponibilità"),
        ("Entro 15 giorni", "prima nota al livello regionale/nazionale: copertura, rischi e decisioni necessarie"),
    ]:
        add_callout(doc, title, body)


def page_sources(doc):
    add_heading(doc, "Luca, questo foglio tienitelo per te", kicker="Nota personale · pagina separata")
    add_p(
        doc,
        "Qui ti lascio, fuori dal documento politico, da dove ho preso i numeri e come leggerli. "
        "Non devi spiegare tutte le fonti durante la riunione: servono a te se qualcuno fa una domanda "
        "o vuole capire quanto lavoro c'è dietro.",
        size=11.2,
        color=NAVY,
        bold=True,
        line=1.24,
        after=10,
    )
    add_callout(
        doc,
        "Come usarlo martedì",
        "Fai il tuo intervento senza leggere il dossier. Quando hai finito il pippone, senza troppa teatralità, "
        "puoi dire: «Ho messo tutto per iscritto, con tempi, dati e fonti». Poi ne lasci una copia a chi decide. "
        "Per quella copia stampa le pagine 1-14; questa pagina 15 resta a te.",
        fill="FFF7D9",
        accent=GOLD,
    )
    add_heading(doc, "Da dove arrivano i dati", level=2)
    sources = [
        ("Futuro Nazionale · Statuto", "https://futuronazionale.it/statuto/"),
        ("Futuro Nazionale · Manifesto V.I.T.A.L.E.", "https://futuronazionale.it/manifesto/"),
        ("Futuro Nazionale · Programma B.I.P. 2027–2037", "https://futuronazionale.it/programma/"),
        ("Futuro Nazionale · Comitati Costituenti", "https://futuronazionale.it/comitati/"),
        ("Provincia MB · Amministrative 2021", "https://www.provincia.mb.it/novita/archivio/Elezioni-Amministrative-2021-00001/"),
        ("Provincia MB · Amministrative 2022", "https://www.provincia.mb.it/scopri-il-territorio-mb/elezioni/elezioni-e-referendum/amministrative/anno-2022/"),
        ("Prefettura MB · scioglimento Bovisio-Masciago", "https://prefettura.interno.gov.it/it/prefetture/monza-e-brianza/comunicati-stampa/avviato-scioglimento-consiglio-comunale-bovisio-masciago"),
        ("Eligendo / DAIT · Europee 2024", "https://elezioni.interno.gov.it/europee/scrutini/20240609/scrutiniEI"),
        ("onData · estrazione machine-readable dei dati Eligendo", "https://github.com/ondata/elezioni_europee_2024"),
        ("Il Giorno · comuni MB al voto nel 2027, 28 agosto 2026", "https://www.ilgiorno.it/monza-brianza/cronaca/elezioni-comunali-brianza-ju2nl2jx"),
        ("Lega · organizzazione e regolamenti lombardi", "https://legaonline.it/lombardia/"),
        ("Patto per il Nord · organigramma e consulte", "https://www.pattoperilnord.it/organigramma"),
    ]
    for label, url in sources:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1.4)
        set_run(p.add_run(label + " · "), size=7.9, color=NAVY, bold=True)
        set_run(p.add_run(url), size=7.35, color=BLUE)
    add_heading(doc, "Come leggere le schede dei comuni", level=2)
    for text in [
        "La percentuale grande è quella ottenuta dal sindaco vincente nell’ultima elezione comunale disponibile: 2021, 2022 oppure 2024 nel caso di Bovisio-Masciago.",
        "Blu significa amministrazione ricondotta al centrodestra, rosso al centrosinistra, verde acqua civica; Bovisio è evidenziata come commissariata.",
        "Nelle tre caselle in basso trovi Europee 2024: area CDX = Fratelli d’Italia + Lega + Forza Italia; area CSX = Partito Democratico + Alleanza Verdi e Sinistra; la terza cifra è l’affluenza.",
        "Quelle europee sono fotografie del territorio, non una previsione di voto per Futuro Nazionale e nemmeno una stima automatica delle future comunali.",
        "Maggio 2027 è un’ipotesi di lavoro: la data ufficiale arriverà soltanto con la convocazione delle elezioni.",
    ]:
        add_bullet(doc, text)
    add_callout(
        doc,
        "Una cosa importante",
        "Le deleghe, le cinque aree e gli obiettivi sono una proposta da discutere e adattare. Presentali come "
        "un metodo aperto alla squadra, non come decisioni già prese.",
        fill=BLUE_LIGHT,
        accent=BLUE,
    )


def configure_private_section(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.header_distance = Cm(.7)
    section.footer_distance = Cm(.7)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(hp.add_run("NOTA PERSONALE · NON CONSEGNARE CON IL DOSSIER"), size=7.2, color=RED, bold=True)
    fp = section.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(fp.add_run("PER LUCA · PROMEMORIA RISERVATO · PAGINA 15"), size=7.0, color=MUTED)


def configure_document(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.55)
    sec.bottom_margin = Cm(1.5)
    sec.left_margin = Cm(1.6)
    sec.right_margin = Cm(1.6)
    sec.header_distance = Cm(.7)
    sec.footer_distance = Cm(.7)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    for level, size, before, after in [(1, 20, 12, 7), (2, 14, 10, 5), (3, 11.5, 7, 3)]:
        st = styles[f"Heading {level}"]
        st.font.name = "Aptos Display"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = rgb(NAVY if level < 3 else BLUE)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    for name in ["List Bullet", "List Bullet 2", "List Number"]:
        st = styles[name]
        st.font.name = "Aptos"
        st.font.size = Pt(9.7)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.15
    # Quiet running furniture.
    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(hp.add_run("FUTURO NAZIONALE · MONZA E BRIANZA"), size=7.2, color=MUTED, bold=True)
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(fp.add_run("PROPOSTA DI LUCA VIVIANI  ·  BOZZA RISERVATA  ·  28.08.2026"), size=7.0, color=MUTED)


def audit(doc):
    assert len(doc.sections) == 2
    sec = doc.sections[0]
    assert round(sec.page_width.cm, 1) == 21.0
    assert round(sec.page_height.cm, 1) == 29.7
    for table in doc.tables:
        assert table.autofit is False
    assert all(name in doc.styles for name in ["Normal", "Heading 1", "Heading 2", "List Bullet", "List Number"])


def build():
    logo = crop_logo()
    eu_data, aggregate = read_eu_data()
    party_chart = chart_parties(aggregate)
    admin_chart = chart_admin()
    doc = Document()
    configure_document(doc)
    page_cover(doc, logo)
    add_page_break(doc)
    page_executive(doc)
    add_page_break(doc)
    page_alignment(doc)
    add_page_break(doc)
    page_org(doc)
    add_page_break(doc)
    page_areas(doc)
    add_page_break(doc)
    page_timeline(doc)
    add_page_break(doc)
    page_political(doc)
    add_page_break(doc)
    page_governance(doc)
    add_page_break(doc)
    page_scenario(doc, admin_chart, party_chart)
    add_page_break(doc)
    pages_municipalities(doc, eu_data)
    add_page_break(doc)
    page_strategy(doc, eu_data)
    add_page_break(doc)
    page_pitch(doc)
    private_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_private_section(private_section)
    page_sources(doc)
    audit(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
