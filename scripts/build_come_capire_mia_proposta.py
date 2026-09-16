from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import build_luca_programma as B  # noqa: E402
from build_guida_discorsiva_luca import TOWNS  # noqa: E402


OUT = ROOT / "output"
DOCX_PATH = OUT / "Numeri_Struttura_Obiettivi_Luca_Viviani.docx"

NAVY = B.NAVY
BLUE = B.BLUE
BLUE_LIGHT = B.BLUE_LIGHT
RED = B.RED
GOLD = B.GOLD
INK = B.INK
MUTED = B.MUTED
LIGHT = B.LIGHT
WHITE = B.WHITE


def para(doc, text, size=10.65, color=INK, bold=False, italic=False,
         after=7, line=1.24, indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.55)
    B.set_run(p.add_run(text), size=size, color=color, bold=bold, italic=italic)
    return p


def lead(doc, text):
    return para(doc, text, size=11.7, color=NAVY, bold=True,
                after=11, line=1.27, indent=False)


def title(doc, page_no, heading):
    B.add_heading(doc, heading, kicker=f"PAGINA {page_no} DEL DOSSIER · SPIEGAZIONE DISCORSIVA")


def takeaway(doc, text, label="IL PUNTO DA FAR CAPIRE"):
    B.add_callout(doc, label, text, fill=BLUE_LIGHT, accent=BLUE)


def page_cover(doc, logo):
    table = doc.add_table(rows=1, cols=2)
    B.set_table_fixed(table, [12.7, 5.1])
    for cell in table.rows[0].cells:
        B.shade(cell, NAVY)
        B.borders(cell, NAVY)
        B.cell_margins(cell, top=180, bottom=180, start=180, end=180)
    left, right = table.rows[0].cells
    left.text = ""
    B.set_run(left.paragraphs[0].add_run("FUTURO NAZIONALE · MONZA E BRIANZA"),
              size=9.2, color=GOLD, bold=True)
    p = left.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    B.set_run(p.add_run("DOCUMENTO DI ACCOMPAGNAMENTO ALLA PROPOSTA"),
              size=8.2, color=WHITE, bold=True)
    right.text = ""
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.add_run().add_picture(str(logo), width=Cm(2.7))

    B.add_p(doc, "NUMERI, STRUTTURA", size=31, color=NAVY, bold=True, before=72, after=0)
    B.add_p(doc, "E OBIETTIVI", size=42, color=RED, bold=True, after=13)
    B.add_p(doc, "LUCA VIVIANI", size=25, color=NAVY, bold=True, after=19)
    B.add_p(doc, "La proposta di coordinamento provinciale di Luca Viviani spiegata pagina per pagina, con parole semplici e in forma discorsiva.",
            size=14.2, color=INK, bold=True, line=1.28, after=17)
    B.add_p(doc, "Un testo per preparare l’intervento, accompagnare la lettura del dossier e illustrare con chiarezza numeri, struttura, tempi e obiettivi.",
            size=11.4, color=MUTED, line=1.27, after=38)
    B.add_metric_strip(doc, [("55", "comuni da organizzare"), ("18", "comuni al voto"),
                             ("100", "primi giorni"), ("1", "squadra provinciale")])
    B.add_p(doc, "DOCUMENTO PER IL CONFRONTO INTERNO", size=8.2, color=MUTED,
            bold=True, before=42, after=2)
    B.add_p(doc, "29 agosto 2026 · Collegato al dossier “Brianza 2027”", size=8.2, color=MUTED)


def page_how_to(doc):
    B.add_heading(doc, "Come leggere questo documento", kicker="PRIMA DI INIZIARE")
    lead(doc, "Il dossier principale presenta rapidamente dati, schemi, funzioni e scadenze. Le pagine che seguono ne sviluppano il significato con un linguaggio discorsivo, collegando ogni elemento al progetto complessivo.")
    para(doc, "Il testo può essere utilizzato in tre modi: come preparazione prima della riunione; come accompagnamento al dossier, seguendo la spiegazione corrispondente a ciascuna pagina; oppure come documento di approfondimento da consegnare a chi desidera comprendere meglio la proposta.")
    para(doc, "Non è necessario memorizzare tutte le percentuali. È invece importante comprenderne il significato. I numeri non promettono voti: dimostrano conoscenza del territorio, distinguono le elezioni comunali da quelle europee e aiutano a prendere decisioni dopo avere raccolto dati, persone e problemi reali.")
    para(doc, "Allo stesso modo, l’organigramma non è un elenco di posti già distribuiti. È una mappa del lavoro da svolgere. Prima vengono le funzioni, gli obiettivi e le regole; solo dopo, attraverso l’ascolto e nel rispetto dello Statuto, vengono individuate le persone adatte. Questa distinzione è essenziale per presentare una proposta inclusiva senza apparire debole o improvvisata.")
    para(doc, "Il tono generale della presentazione deve essere fermo ma non arrogante. Il progetto dimostra iniziativa e capacità di guida, senza sostenere che ogni dettaglio sia già deciso. La richiesta è quella di un mandato per costruire insieme, rendere il lavoro verificabile e consegnare risultati entro tempi precisi.")
    takeaway(doc, "Il dossier dimostra preparazione. Questo testo rende visibile la logica che tiene insieme tutte le sue pagine.")


def page_1(doc):
    title(doc, 1, "Perché si parte da 55 comuni, 18 scadenze e 100 giorni")
    lead(doc, "La copertina non contiene soltanto uno slogan. Riassume immediatamente la dimensione del compito proposto: organizzare una provincia intera, partire dalle scadenze più vicine e trasformare il tempo disponibile in lavoro concreto.")
    para(doc, "Il riferimento ai cinquantacinque comuni chiarisce che il coordinamento provinciale non può vivere soltanto nei luoghi dove esistono già amicizie o contatti. Deve costruire una presenza capace di raggiungere progressivamente tutta la provincia. Nessuna persona può seguire da sola cinquantacinque realtà: per questo la proposta si fonda sulla squadra, sulle deleghe e sulle aree operative.")
    para(doc, "I diciotto comuni indicati sono quelli interessati dalla prossima tornata amministrativa sulla base delle scadenze oggi disponibili. Sono la priorità perché l’appuntamento elettorale obbliga a lavorare subito. Non significa dimenticare gli altri trentasette comuni; significa usare bene il tempo e concentrare inizialmente le energie dove una struttura organizzata sarà necessaria prima.")
    para(doc, "La cifra di circa quattrocentotrentacinquemila cittadini descrive la popolazione complessivamente coinvolta nei diciotto comuni. Non è una previsione degli elettori di Futuro Nazionale. Serve a mostrare che la sfida non è marginale: riguarda una parte molto rilevante della Brianza, compreso il capoluogo e numerosi centri importanti.")
    para(doc, "La parola chiave è preparazione. Il movimento non deve comparire poche settimane prima del voto con qualche manifesto e una lista costruita in fretta. Deve iniziare ora a conoscere persone, problemi, amministratori, associazioni e possibili interlocutori. Il senso di ‘Organizzare, Radicare, Preparare’ è esattamente questo.")
    takeaway(doc, "Non viene richiesto un incarico generico, ma la responsabilità di trasformare una presenza ancora incompleta in una struttura provinciale pronta per il 2027.")


def page_2(doc):
    title(doc, 2, "La sintesi: che cosa viene proposto e perché Luca può guidarlo")
    lead(doc, "La seconda pagina risponde subito a due domande: quale risultato si vuole ottenere e quali caratteristiche rendono Luca Viviani adatto a guidare questo lavoro.")
    para(doc, "La proposta è un coordinamento provinciale che non cancelli ciò che esiste, ma lo colleghi. Comitati, aderenti, amministratori e persone disponibili devono riconoscersi in un unico metodo di lavoro. Entro cento giorni dovranno essere prodotti una mappa completa, un assetto operativo, cinque aree territoriali e diciotto dossier comunali. Questi risultati permetteranno di verificare il lavoro senza affidarsi soltanto alle impressioni.")
    para(doc, "L’esperienza istituzionale di Luca è utile perché ha ricoperto ruoli di consigliere provinciale e assessore comunale. Questi incarichi non vengono ricordati per celebrare il passato, ma per dimostrare conoscenza degli enti locali, dei loro tempi, delle competenze e delle difficoltà che separano una dichiarazione politica da una decisione realmente applicabile.")
    para(doc, "Luca è architetto e ha lavorato per molti anni come docente negli istituti superiori. L’architettura gli ha dato capacità di lettura del territorio, dei servizi, delle infrastrutture e dei vincoli. L’insegnamento gli ha dato metodo, capacità di spiegare e abitudine a far crescere le persone. Oggi, essendo pensionato, può dedicare al coordinamento anche le ore diurne, seguendo incontri e problemi con continuità.")
    para(doc, "Il quarto elemento è il modello di leadership proposto. Le attività non devono concentrarsi su una sola persona: occorre assegnare responsabilità vere, con un risultato atteso e una verifica. Una squadra larga è utile soltanto se ciascuno sa che cosa deve fare e se il coordinatore mantiene la direzione complessiva.")
    takeaway(doc, "Luca chiede di essere valutato non soltanto per il curriculum, ma per quattro risorse unite insieme: esperienza, competenza, tempo operativo e capacità di costruire una squadra.")


def page_3(doc):
    title(doc, 3, "Il mandato: regole chiare e identità riconoscibile")
    lead(doc, "Questa pagina chiarisce che la proposta nasce dentro le regole del movimento. La funzione corretta è quella di Coordinatore Provinciale e la sua attribuzione deve seguire le procedure previste dallo Statuto.")
    para(doc, "Non viene proposta una struttura personale né un’organizzazione parallela. Si richiede il mandato necessario per guidare il lavoro provinciale, mentre nomine statutarie, candidature, accordi elettorali e decisioni di maggiore rilievo restano affidati agli organi competenti. L’autonomia deve essere sufficiente per lavorare, ma sempre tracciabile e coerente con la linea regionale e nazionale.")
    para(doc, "Nel dossier traduco i valori V.I.T.A.L.E. in comportamenti. Virtù significa mantenere la parola data, rispettare le scadenze e rendere conto. Identità significa dare alla Brianza una linea riconoscibile senza staccarla dal movimento nazionale. Tradizioni significa essere presenti nei paesi e nelle comunità reali. Amore significa prendersi cura delle famiglie e delle fragilità. Libertà significa delegare iniziativa dentro regole chiare. Eccellenza significa scegliere per competenza e correggere ciò che non funziona.")
    para(doc, "Il principio di fondo è avere una sola direzione politica provinciale e molte responsabilità operative distribuite. I comitati esistenti non devono essere scavalcati o cancellati: devono essere ascoltati, riconosciuti e messi in condizione di collaborare con standard comuni. In questo modo l’unità non diventa centralismo e l’inclusione non diventa confusione.")
    takeaway(doc, "La richiesta riguarda l’autorità necessaria per costruire, non la libertà di agire senza regole. La forza del coordinamento deve venire dall’unione tra identità, disciplina e capacità operativa.")


def page_4(doc):
    title(doc, 4, "L’organizzazione: una squadra larga con responsabilità nette")
    lead(doc, "L’organigramma non distribuisce incarichi in anticipo. Descrive il lavoro che una provincia organizzata deve essere capace di svolgere e distingue con chiarezza la direzione politica dalle funzioni operative.")
    para(doc, "Il Coordinatore Provinciale mantiene indirizzo politico, rappresentanza e responsabilità complessiva. L’Esecutivo Provinciale, quando costituito secondo lo Statuto, traduce l’indirizzo in decisioni e verifica il lavoro. Le funzioni di organizzazione, tesseramento, disciplina, giovani e segreteria garantiscono continuità: eventi e calendario, qualità dei dati, rispetto delle regole, formazione della nuova classe dirigente, verbali e controllo delle scadenze.")
    para(doc, "Accanto agli organi statutari vengono proposte deleghe operative. Servono referenti per le aree territoriali, per gli enti locali, per le amministrative del 2027, per programma e consulte, comunicazione, eventi, dati, monitoraggio, conformità e privacy. Queste deleghe non creano piccoli centri di potere: rendono visibile chi deve produrre un determinato risultato.")
    para(doc, "I nomi non vengono indicati perché prima bisogna ascoltare e conoscere le disponibilità. Partire dai nomi farebbe apparire il progetto come una spartizione. Partendo dalle funzioni, invece, è possibile chiedere a ogni persona quale responsabilità sia pronta ad assumere, con quali competenze e con quale disponibilità concreta.")
    para(doc, "Ogni funzione deve avere un responsabile unico, un obiettivo scritto, indicatori semplici e una verifica mensile. Questo permette di coinvolgere molte persone senza perdere il controllo e consente anche di correggere rapidamente un incarico che non produce risultati.")
    takeaway(doc, "Prima definiamo il lavoro da fare; poi scegliamo le persone. Nessuno viene escluso in partenza, ma nessun ruolo deve essere soltanto simbolico.")


def page_5(doc):
    title(doc, 5, "Le cinque aree: avvicinare il coordinamento ai comuni")
    lead(doc, "Le cinque aree operative servono a rendere gestibile una provincia composta da cinquantacinque comuni. Non sono nuovi livelli statutari e non dividono la Brianza in cinque feudi: distribuiscono il carico e riducono le distanze.")
    para(doc, "La prima area comprende Monza e la cintura vicina; la seconda segue l’Ovest e l’asse del Seveso; la terza la Brianza centrale; la quarta la Valle del Lambro; la quinta il Vimercatese. Il perimetro iniziale tiene conto soprattutto dei diciotto comuni al voto e potrà essere corretto dopo avere censito tutti i comuni, i referenti e le realtà già attive.")
    para(doc, "Ogni area deve garantire pochi standard, ma verificabili: una riunione operativa mensile, un calendario condiviso, un contatto affidabile per ogni comune, un incontro territoriale di ascolto e un breve rendiconto. Dove esistono le condizioni previste, deve inoltre accompagnare la nascita o il riconoscimento dei comitati.")
    para(doc, "Il referente d’area non decide la linea politica e non diventa il proprietario dei comuni. Deve facilitare i rapporti, far circolare informazioni, sostenere i referenti locali e segnalare problemi alla struttura provinciale. La direzione resta unica, mentre la presenza diventa più vicina e continua.")
    para(doc, "Questa scelta è particolarmente importante nei comuni più piccoli, dove spesso non ci sono abbastanza persone per mantenere da sole tutte le attività. Comuni vicini possono condividere incontri, formazione e competenze, conservando però una lettura specifica dei propri problemi.")
    takeaway(doc, "Cinque aree operative, ma una sola provincia: messaggio, regole, dati e responsabilità restano comuni.")


def page_6(doc):
    title(doc, 6, "Il cronoprogramma: che cosa accade dopo la nomina")
    lead(doc, "I cento giorni servono a trasformare il mandato in risultati osservabili. Non sono uno slogan: sono una sequenza di lavoro che parte dall’ascolto e arriva a una macchina organizzativa pronta.")
    para(doc, "Nei primi quindici giorni la priorità sarà incontrare comitati, referenti, amministratori e persone disponibili. Prima di assegnare ruoli occorre conoscere ciò che esiste, le attività già svolte, le competenze, le criticità e anche le eventuali tensioni. Partire con l’ascolto significa riconoscere il lavoro degli altri e ridurre il rischio di decisioni basate soltanto sulle relazioni personali.")
    para(doc, "Entro trenta giorni dovrà essere completato il censimento dei cinquantacinque comuni: dove il movimento è presente, dove è assente, quali amministratori o aderenti esistono, quali competenze possono essere utilizzate e quali dati devono essere aggiornati. Questa fotografia è la base per organizzare il resto.")
    para(doc, "Entro sessanta giorni devono essere operative le funzioni principali, le cinque aree, il calendario provinciale e le regole di lavoro. Entro novanta giorni dovranno essere disponibili diciotto dossier comunali con situazione politica, problemi, relazioni, possibili formule elettorali, fabbisogni e rischi. Da gennaio a maggio 2027 il lavoro passerà gradualmente dalla costruzione alle decisioni e poi all’esecuzione della campagna.")
    para(doc, "Il cruscotto mensile serve a non raccontarci che tutto procede bene senza prove. Controlleremo copertura dei comuni, iscritti e rinnovi, attività realizzate, dossier completati, scadenze non rispettate e qualità del seguito dato alle decisioni. Se qualcosa si blocca, deve emergere in tempo per essere corretto.")
    takeaway(doc, "Dopo cento giorni non si promettono cinquantacinque sezioni perfette, ma una mappa affidabile, una struttura leggibile e un piano concreto per colmare i vuoti.")


def page_7(doc):
    title(doc, 7, "Gli otto laboratori: trasformare i temi in proposte brianzole")
    lead(doc, "Una struttura politica non può vivere soltanto di riunioni, ruoli e banchetti. Deve produrre contenuti. Gli otto laboratori servono a tradurre il programma nazionale in proposte locali serie e verificabili.")
    para(doc, "I temi sono sicurezza e legalità; impresa e lavoro; famiglia e comunità; scuola e giovani; territorio e ambiente; mobilità e infrastrutture; identità e cultura; salute e sociale. Sono ambiti che riguardano direttamente la vita brianzola e permettono di coinvolgere amministratori, professionisti, associazioni e cittadini sulla base delle competenze.")
    para(doc, "Ogni laboratorio deve partire da un problema osservabile, non da uno slogan. Deve capire se la competenza appartiene al Comune, alla Provincia, alla Regione o allo Stato; consultare più fonti; distinguere il principio politico dalla soluzione tecnica; indicare vincoli e costi; pubblicare soltanto ciò che ha ricevuto l’approvazione necessaria.")
    para(doc, "Entro febbraio 2027 ogni laboratorio dovrebbe produrre almeno una scheda breve. Lo scopo non è riempire pagine, ma consegnare ai territori materiali utili per programmi, incontri e confronto pubblico. In questo modo una consulta non diventa un premio o un titolo, ma una fabbrica di proposte e una palestra per futura classe dirigente.")
    para(doc, "Questa parte del progetto è importante anche per selezionare le persone. Chi sa ascoltare, studiare, scrivere e trasformare un problema in una proposta realistica dimostra sul campo di poter assumere responsabilità maggiori.")
    takeaway(doc, "Le consulte non devono distribuire visibilità: devono produrre idee utilizzabili, competenze e persone credibili.")


def page_8(doc):
    title(doc, 8, "Le regole di governo: delegare senza perdere direzione")
    lead(doc, "Una squadra funziona quando tutti sanno chi può decidere, entro quali limiti e quale prova deve lasciare. Le regole servono a proteggere insieme velocità, unità politica e reputazione.")
    para(doc, "Le decisioni operative e reversibili possono essere delegate rapidamente, purché rientrino nel mandato, nel calendario e nelle risorse approvate. Le comunicazioni pubbliche provinciali richiedono invece una linea condivisa e l’approvazione del coordinatore. Accordi, candidature e alleanze non possono essere promessi autonomamente: devono essere portati agli organi competenti.")
    para(doc, "Anche spese, raccolta fondi e utilizzo dei dati personali richiedono procedure chiare e tracciabili. I dati devono essere ridotti a quelli necessari, accessibili soltanto a chi ne ha bisogno e separati quando hanno natura diversa. Gli incarichi devono poggiare su competenza, obiettivi, durata e dichiarazione di eventuali conflitti d’interesse.")
    para(doc, "La linea etica è netta: le responsabilità interne non sono una moneta con cui promettere candidature, nomine o posti in enti e società. Chi lavora nella struttura lo fa per produrre risultati politici e organizzativi. Ogni eventuale proposta futura deve essere valutata dall’organo competente sulla base di capacità, trasparenza e assenza di conflitti.")
    para(doc, "Queste regole non indeboliscono il coordinatore. Al contrario, gli consentono di delegare senza perdere controllo e proteggono il movimento da decisioni personali difficili da correggere.")
    takeaway(doc, "L’obiettivo è una squadra veloce nelle attività quotidiane e prudente nelle decisioni che possono produrre conseguenze politiche, economiche o reputazionali.")


def page_9(doc):
    title(doc, 9, "Lo scenario 2027: come leggere il campo di gioco")
    lead(doc, "La pagina elettorale non pretende di prevedere il risultato di Futuro Nazionale. Descrive il terreno sul quale dobbiamo lavorare e mette insieme, con molta cautela, le ultime amministrative e le Europee del 2024.")
    para(doc, "I diciotto comuni comprendono dieci centri sopra i quindicimila abitanti e otto sotto. Nell’ultima situazione disponibile, dodici avevano visto una vittoria del centrodestra, cinque del centrosinistra e uno di una lista civica; Bovisio-Masciago è successivamente passato al commissariamento. Questo dato fotografa il passato amministrativo, non stabilisce il futuro.")
    para(doc, "Il 50,8 per cento indicato come area di centrodestra alle Europee è la somma di Fratelli d’Italia, Lega e Forza Italia-Noi Moderati. L’area di centrosinistra utilizzata nel confronto somma Partito Democratico e Alleanza Verdi e Sinistra. Sono aggregazioni analitiche, non coalizioni già definite per le comunali, e non comprendono tutti i partiti presenti.")
    para(doc, "Futuro Nazionale non era sulla scheda delle Europee 2024. Non è quindi possibile sostenere che una parte di quei voti gli appartenga automaticamente. Il dato può però essere usato come indicatore del clima politico: aiuta a capire dove il centrodestra ha un bacino forte, dove la competizione è più equilibrata e dove serve un’offerta locale particolarmente autonoma e credibile.")
    para(doc, "La differenza tra amministrative ed Europee è già un insegnamento. Monza, Cesano Maderno, Lesmo e Verano sono amministrate dal centrosinistra pur avendo mostrato un vantaggio europeo dell’area di centrodestra. Significa che candidato, coalizione, liste civiche, reputazione e problemi locali possono contare più dell’appartenenza nazionale.")
    takeaway(doc, "I numeri non dicono quanti voti otterrà il movimento. Indicano quali domande porre prima di scegliere candidati, alleanze, priorità e investimenti organizzativi.")


def town_text(doc, name):
    orientation, municipal, cdx, csx, turnout, reading = TOWNS[name]
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    B.set_run(p.add_run(name), size=11.4, color=NAVY, bold=True)
    para(doc, f"La lettura parte dall’ultima situazione comunale: {orientation}, con il risultato del vincitore al {municipal}. Alle Europee 2024 l’area CDX è al {cdx}, l’area CSX al {csx} e l’affluenza al {turnout}. {reading}",
         size=9.75, after=6, line=1.17, indent=False)


def towns_page(doc, dossier_page, names, part):
    title(doc, dossier_page, f"I comuni: leggere le schede senza farsi ingannare · {part}/3")
    lead(doc, "Per ogni comune vengono accostate tre informazioni: chi ha vinto l’ultima amministrativa, come si sono distribuite le principali aree alle Europee 2024 e quante persone hanno votato. Il confronto serve a formulare domande, non ad assegnare voti futuri.")
    for name in names:
        town_text(doc, name)


def page_13(doc):
    title(doc, 13, "Dai numeri alle priorità organizzative")
    lead(doc, "Dopo avere osservato i diciotto comuni, li raggruppo in tre fasce. Non sono fasce di consenso per Futuro Nazionale e non sono graduatorie di comuni facili o difficili: indicano il tipo di lavoro che appare necessario sulla base del solo benchmark europeo.")
    para(doc, "Nei bacini dove l’area di centrodestra è molto forte, come Briosco, Meda, Verano, Lentate, Seveso, Biassono, Lesmo e Cesano Maderno, esiste un pubblico potenzialmente sensibile ai nostri temi. Esistono però anche partiti, amministratori e reti già radicati. Dobbiamo quindi dimostrare un valore aggiunto attraverso persone, proposte e capacità organizzativa.")
    para(doc, "Nell’area competitiva - Lissone, Limbiate, Bovisio-Masciago, Varedo, Sulbiate, Vedano al Lambro e Arcore - la presenza continuativa, le alleanze locali e la qualità della proposta possono incidere molto. Qui non basta richiamare un’identità nazionale: bisogna costruire credibilità comune per comune.")
    para(doc, "Monza, Carnate e Vimercate formano il contesto più complesso del gruppo. Richiedono lavoro anticipato, reputazione locale e un’offerta capace di distinguersi. In particolare, Vimercate è l’unico comune dei diciotto dove il benchmark PD più AVS supera, anche se di poco, la somma dei tre partiti dell’area di centrodestra considerati.")
    para(doc, "Da qui derivano cinque priorità entro dicembre: unificare ciò che esiste, coprire ciascuno dei diciotto comuni con un referente operativo, conoscere storia e problemi locali attraverso dossier, formare portavoce e organizzatori, portare ai livelli competenti decisioni chiare con rischi e alternative.")
    takeaway(doc, "La priorità non si decide guardando soltanto una percentuale. Si decide unendo dato elettorale, forza organizzativa reale, persone disponibili, problemi locali e possibilità di costruire una proposta credibile.")


def page_14(doc):
    title(doc, 14, "La richiesta finale: un mandato da verificare")
    lead(doc, "La proposta si chiude tornando al punto iniziale: non viene richiesto un titolo da esibire, ma un mandato di costruzione accompagnato da scadenze, responsabilità e controlli.")
    para(doc, "Chiedo il riconoscimento del ruolo di Coordinatore Provinciale secondo le procedure transitorie previste. Chiedo novanta giorni per completare censimento, assetto operativo e dossier dei diciotto comuni, con una prima verifica a trenta giorni e una seconda a novanta. Chiedo la possibilità di proporre deleghe e referenti d’area senza anticipare nomi e rispettando gli organi statutari. Chiedo infine l’accesso ai dati e ai contatti strettamente necessari, nel rispetto della privacy.")
    para(doc, "In caso di conferimento del mandato, entro quarantotto ore partirà un messaggio unitario ai comitati e ai referenti, fondato su ascolto, continuità e calendario. Entro sette giorni sarà convocata la prima riunione operativa con una mappa preliminare delle presenze e delle disponibilità. Entro quindici giorni sarà prodotta una prima nota per il livello regionale e nazionale con copertura, rischi e decisioni necessarie.")
    para(doc, "Il punto politico è semplice: la Brianza non ha bisogno di un coordinatore che occupi uno spazio, ma di una guida che costruisca una squadra. L’esperienza, le competenze e la disponibilità di tempo di Luca hanno valore soltanto se si trasformano in risultati verificabili e se consentono ad altre persone capaci di assumere responsabilità reali.")
    para(doc, "Presento questa struttura come una base di lavoro aperta. Non sostengo che ogni dettaglio sia già definito. Sostengo però che il metodo, i tempi e gli obiettivi debbano essere chiari fin dall’inizio, perché l’apertura senza direzione produce immobilismo e la direzione senza ascolto produce divisioni.")
    takeaway(doc, "La richiesta è un mandato verificabile: entro cento giorni sarà possibile vedere che cosa è stato costruito, dove il movimento è presente e quali decisioni servono per arrivare pronti al 2027.")


def speech_1(doc):
    B.add_heading(doc, "Il discorso completo · apertura", kicker="VERSIONE PRONUNCIABILE")
    lead(doc, "Le pagine seguenti uniscono l’intero dossier in un solo intervento. Posso leggerle, ridurle oppure usarle come traccia, mantenendo la stessa logica.")
    para(doc, "Ho preparato questa proposta perché credo che Monza e Brianza abbia bisogno di passare rapidamente dalla presenza spontanea a una struttura provinciale riconoscibile. Non partiamo dal nulla: ci sono persone, comitati, amministratori, esperienze e disponibilità. Quello che manca è un metodo comune che colleghi queste energie, distribuisca il lavoro e permetta di verificare i risultati.", indent=False)
    para(doc, "La provincia comprende cinquantacinque comuni e diciotto di questi sono interessati dalla prossima tornata amministrativa. Parliamo di territori che coinvolgono complessivamente circa quattrocentotrentacinquemila cittadini. Non sono numeri messi in copertina per fare scena. Servono a far capire la dimensione della responsabilità: una persona da sola non può seguire tutto e una campagna improvvisata poche settimane prima del voto non sarebbe sufficiente.", indent=False)
    para(doc, "Per questo propongo di lavorare fin da subito su tre verbi: organizzare, radicare e preparare. Organizzare significa sapere chi fa che cosa. Radicare significa essere presenti nei comuni anche lontano dalle elezioni. Preparare significa arrivare alle decisioni del 2027 con dati, persone, proposte e alternative già studiate.", indent=False)
    para(doc, "Metto a disposizione l’esperienza maturata come consigliere provinciale e assessore comunale, la mia formazione di architetto, molti anni di insegnamento negli istituti superiori e soprattutto la disponibilità di tempo che oggi posso garantire. Non presento il mio curriculum come un elenco di titoli: lo presento come un insieme di strumenti utili per conoscere gli enti locali, leggere il territorio, spiegare questioni complesse e accompagnare una squadra.", indent=False)
    takeaway(doc, "Il tono dell’apertura deve essere concreto: riconosco ciò che esiste, descrivo il problema e spiego quali risorse metto a disposizione.", label="CHIAVE DI LETTURA")


def speech_2(doc):
    B.add_heading(doc, "Il discorso completo · organizzazione", kicker="VERSIONE PRONUNCIABILE")
    para(doc, "Il modello che propongo distingue gli organi previsti dallo Statuto dalle deleghe operative. Il Coordinatore Provinciale mantiene la direzione politica e la responsabilità complessiva; intorno a lui devono funzionare l’Esecutivo, l’organizzazione, il tesseramento, la disciplina, il settore giovani e una segreteria capace di seguire agenda, verbali e scadenze. Accanto a queste funzioni servono responsabilità operative su territori, enti locali, amministrative, programma, comunicazione, eventi, dati e privacy.", indent=False)
    para(doc, "Non ho inserito nomi perché non voglio presentare una squadra già spartita prima dell’ascolto. Prima definiamo il lavoro, poi individuiamo le persone più adatte. L’inclusione, per me, non consiste nel moltiplicare i titoli: consiste nel dare a più persone la possibilità di assumere un compito vero, con un obiettivo e una verifica.", indent=False)
    para(doc, "Per seguire cinquantacinque comuni propongo cinque aree operative: Monza e cintura, Ovest e asse del Seveso, Brianza centrale, Valle del Lambro e Vimercatese. Non sono strutture autonome e non diventano correnti. Servono a ridurre le distanze, garantire riunioni regolari, raccogliere problemi e mantenere un contatto verificato nei comuni.", indent=False)
    para(doc, "Nei primi quindici giorni ascolterò le realtà esistenti. Entro trenta completerò la mappa della provincia. Entro sessanta renderò operative le funzioni e le cinque aree. Entro novanta preparerò i dossier dei diciotto comuni al voto. Ogni mese presenterò un rendiconto sintetico su copertura, tesseramento, attività, preparazione, ritardi e qualità del lavoro.", indent=False)
    para(doc, "Parallelamente voglio attivare otto laboratori sui temi che riguardano la Brianza: sicurezza, impresa e lavoro, famiglia, scuola e giovani, territorio, mobilità, identità, salute e sociale. Ogni laboratorio dovrà produrre proposte brevi, documentate e compatibili con le competenze degli enti locali. Le consulte devono essere luoghi di lavoro e formazione, non contenitori di incarichi.", indent=False)
    takeaway(doc, "Questa parte dimostra che so trasformare la leadership in un sistema di responsabilità, tempi e risultati.", label="CHIAVE DI LETTURA")


def speech_3(doc):
    B.add_heading(doc, "Il discorso completo · territorio e richiesta", kicker="VERSIONE PRONUNCIABILE")
    para(doc, "Ho analizzato anche il campo elettorale dei diciotto comuni, mettendo a confronto l’ultima amministrazione disponibile con le Europee del 2024. I dati mostrano un’area di centrodestra generalmente forte, ma mostrano anche che le elezioni comunali seguono logiche diverse. Candidato, coalizione, liste civiche, reputazione e temi locali possono cambiare completamente il risultato.", indent=False)
    para(doc, "Futuro Nazionale non era presente sulla scheda europea e quindi quei voti non possono essere considerati nostri. Le percentuali sono un punto di partenza per fare domande: dove esiste un bacino sensibile ai nostri temi? Dove troviamo una concorrenza già organizzata? Dove il centrosinistra è più competitivo? Dove la partecipazione è bassa? Solo unendo questi dati alla forza reale dei referenti, alle relazioni e ai problemi locali potremo decidere come muoverci.", indent=False)
    para(doc, "Per questo non propongo promesse elettorali premature. Propongo cinque azioni: unificare le realtà esistenti, coprire i diciotto comuni, conoscere ogni territorio attraverso un dossier, formare persone capaci e portare agli organi competenti opzioni chiare. Le candidature e le alleanze non si improvvisano e non si promettono individualmente.", indent=False)
    para(doc, "Chiedo quindi il riconoscimento del ruolo di Coordinatore Provinciale secondo le procedure previste, un mandato operativo iniziale di novanta-cento giorni, la possibilità di proporre deleghe e referenti dopo l’ascolto e l’accesso alle informazioni strettamente necessarie. Entro quarantotto ore dalla nomina invierò un messaggio unitario; entro sette giorni terrò la prima riunione operativa; entro quindici presenterò una prima nota sullo stato della provincia.", indent=False)
    para(doc, "La Brianza non ha bisogno di un coordinatore che occupi uno spazio, ma di una guida che costruisca una squadra. Io sono pronto ad assumermi questa responsabilità. Non vi chiedo fiducia al buio: vi propongo un mandato da verificare. Entro cento giorni saprete che cosa è stato costruito, dove siamo presenti, quali problemi restano aperti e quali decisioni servono per arrivare pronti al 2027.", indent=False)
    takeaway(doc, "Chiudere qui, senza aggiungere promesse personali. La proposta termina con una responsabilità misurabile e una richiesta chiara.", label="CHIUSURA")


def configure(doc):
    B.configure_document(doc)
    sec = doc.sections[0]
    header = sec.header
    header.paragraphs[0].text = ""
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    B.set_run(hp.add_run("NUMERI, STRUTTURA E OBIETTIVI · LUCA VIVIANI"),
              size=7.2, color=MUTED, bold=True)
    footer = sec.footer
    footer.paragraphs[0].text = ""
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    B.set_run(fp.add_run("DOCUMENTO DI ACCOMPAGNAMENTO · CONFRONTO INTERNO · 29.08.2026"),
              size=7.0, color=MUTED)
    normal = doc.styles["Normal"]
    normal.paragraph_format.line_spacing = 1.24
    normal.paragraph_format.space_after = Pt(7)


def audit(doc):
    assert len(doc.sections) == 1
    assert round(doc.sections[0].page_width.cm, 1) == 21.0
    assert round(doc.sections[0].page_height.cm, 1) == 29.7
    assert len(doc.inline_shapes) == 1
    for table in doc.tables:
        assert table.autofit is False
        for row in table.rows:
            assert row.height is None
    text = "\n".join(p.text for p in doc.paragraphs).lower()
    required = ["numeri, struttura", "e obiettivi", "io sono pronto", "pagina 14 del dossier",
                "cinquantacinque comuni", "futuro nazionale non era presente",
                "non vi chiedo fiducia al buio"]
    missing = [x for x in required if x not in text]
    assert not missing, missing


def build():
    logo = B.crop_logo()
    doc = Document()
    configure(doc)
    pages = [
        lambda: page_cover(doc, logo),
        lambda: page_how_to(doc),
        lambda: page_1(doc),
        lambda: page_2(doc),
        lambda: page_3(doc),
        lambda: page_4(doc),
        lambda: page_5(doc),
        lambda: page_6(doc),
        lambda: page_7(doc),
        lambda: page_8(doc),
        lambda: page_9(doc),
        lambda: towns_page(doc, 10, ["Arcore", "Biassono", "Bovisio-Masciago", "Briosco", "Carnate", "Cesano Maderno"], 1),
        lambda: towns_page(doc, 11, ["Lentate sul Seveso", "Lesmo", "Limbiate", "Lissone", "Meda", "Monza"], 2),
        lambda: towns_page(doc, 12, ["Seveso", "Sulbiate", "Varedo", "Vedano al Lambro", "Verano Brianza", "Vimercate"], 3),
        lambda: page_13(doc),
        lambda: page_14(doc),
        lambda: speech_1(doc),
        lambda: speech_2(doc),
        lambda: speech_3(doc),
    ]
    for i, fn in enumerate(pages):
        if i:
            B.add_page_break(doc)
        fn()
    audit(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
