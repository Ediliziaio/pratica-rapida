from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import build_luca_programma as B  # noqa: E402


OUT = ROOT / "output"
DOCX_PATH = OUT / "Guida_Discorsiva_Proposta_Luca_Viviani.docx"

NAVY = B.NAVY
BLUE = B.BLUE
BLUE_LIGHT = B.BLUE_LIGHT
RED = B.RED
RED_LIGHT = B.RED_LIGHT
GOLD = B.GOLD
INK = B.INK
MUTED = B.MUTED
LIGHT = B.LIGHT
WHITE = B.WHITE
TEAL = B.TEAL
GRAY = B.GRAY


def para(doc, text, size=10.4, color=INK, bold=False, italic=False, after=7, line=1.22, first=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if first:
        p.paragraph_format.first_line_indent = Cm(0.55)
    B.set_run(p.add_run(text), size=size, color=color, bold=bold, italic=italic)
    return p


def lead(doc, text):
    return para(doc, text, size=11.5, color=NAVY, bold=True, after=10, line=1.25, first=False)


def note(doc, title, body, fill=BLUE_LIGHT, accent=BLUE):
    B.add_callout(doc, title, body, fill=fill, accent=accent)


def section_title(doc, title, kicker):
    B.add_heading(doc, title, kicker=kicker)


def page_cover(doc, logo):
    table = doc.add_table(rows=1, cols=2)
    B.set_table_fixed(table, [12.7, 5.1])
    for cell in table.rows[0].cells:
        B.shade(cell, NAVY)
        B.borders(cell, NAVY)
        B.cell_margins(cell, top=180, bottom=180, start=180, end=180)
    left, right = table.rows[0].cells
    left.text = ""
    p = left.paragraphs[0]
    B.set_run(p.add_run("GUIDA DISCORSIVA ALLA PROPOSTA"), size=9.2, color=GOLD, bold=True)
    p2 = left.add_paragraph()
    p2.paragraph_format.space_before = Pt(4)
    B.set_run(p2.add_run("MONZA E BRIANZA"), size=12, color=WHITE, bold=True)
    right.text = ""
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.add_run().add_picture(str(logo), width=Cm(2.7))

    B.add_p(doc, "CAPIRE", size=36, color=NAVY, bold=True, before=72, after=0)
    B.add_p(doc, "LA PROPOSTA", size=44, color=RED, bold=True, after=10)
    B.add_p(doc, "DI LUCA VIVIANI", size=25, color=NAVY, bold=True, after=18)
    B.add_p(doc, "Spiegazione completa, in parole semplici, del progetto politico e organizzativo per il coordinamento provinciale.", size=14, color=INK, bold=True, line=1.28, after=16)
    B.add_p(doc, "Un testo pensato per chi vuole comprendere il senso dei numeri, delle strutture e delle scadenze senza essere un esperto di organizzazione politica o di analisi elettorale.", size=11.2, color=MUTED, line=1.25, after=36)
    B.add_metric_strip(doc, [("55", "comuni della provincia"), ("18", "al voto nel 2027"),
                             ("100", "giorni per costruire"), ("1", "squadra provinciale")])
    B.add_p(doc, "DOCUMENTO DI ACCOMPAGNAMENTO · USO INTERNO", size=8.2, color=MUTED, bold=True, before=42, after=2)
    B.add_p(doc, "29 agosto 2026 · Basato sul dossier “Brianza 2027”", size=8.2, color=MUTED)


def page_premessa(doc):
    section_title(doc, "A che cosa serve questa guida", "Premessa")
    lead(doc, "Il dossier originale è volutamente sintetico: usa schemi, tabelle, percentuali e parole chiave per essere consultato rapidamente durante una riunione. Questa guida fa il lavoro opposto. Prende gli stessi contenuti e li racconta con calma, come farebbe una persona che li ha preparati e vuole renderli comprensibili a chiunque.")
    para(doc, "Non troverai nuove promesse, nuovi nomi o decisioni aggiunte dopo. Troverai invece il significato delle cose già scritte: perché compaiono cinquantacinque comuni, perché diciotto hanno una priorità immediata, che cosa vuol dire organizzare la provincia in cinque aree, perché si parla di cento giorni e come devono essere interpretati i risultati delle ultime elezioni comunali ed europee.")
    para(doc, "La proposta, infatti, non è un programma elettorale nel senso tradizionale. Non elenca soltanto temi come sicurezza, lavoro, scuola o trasporti. Prima ancora, prova a rispondere a una domanda più concreta: se Futuro Nazionale vuole diventare una presenza reale in Monza e Brianza, chi fa che cosa, in quali tempi, con quali responsabilità e con quali controlli?")
    para(doc, "Questo è il filo che tiene insieme tutto il documento. Luca Viviani non si presenta dicendo semplicemente di avere esperienza o di conoscere molte persone. Si presenta dicendo che quell’esperienza può essere trasformata in una macchina organizzativa leggibile, aperta e verificabile. In altre parole, non chiede fiducia al buio: propone un modo per controllare nel tempo se la fiducia è stata ben riposta.")
    note(doc, "Come leggerla", "Si può leggere dall’inizio alla fine oppure usare i titoli per ritrovare un argomento. Le pagine sui comuni spiegano i dati uno per uno; le ultime pagine chiariscono che cosa Luca chiede realmente alla riunione e come dovrebbe presentare il dossier.", fill="FFF7D9", accent=GOLD)


def page_big_picture(doc):
    section_title(doc, "La proposta, detta in una frase", "Il senso complessivo")
    lead(doc, "L’idea è questa: dare a Luca un mandato per unire ciò che già esiste, distribuire il lavoro tra più persone e arrivare alle amministrative del 2027 con una struttura provinciale riconoscibile, non con una somma di iniziative isolate.")
    para(doc, "La provincia di Monza e Brianza comprende cinquantacinque comuni. Questo numero serve a ricordare la dimensione vera del compito. Un coordinatore provinciale non può limitarsi ai paesi dove conosce già qualcuno, né può immaginare di fare tutto personalmente. Deve costruire un sistema capace, gradualmente, di osservare e servire l’intero territorio.")
    para(doc, "All’interno di questi cinquantacinque comuni, diciotto dovrebbero andare al voto amministrativo nella primavera del 2027. Sono quindi la priorità immediata. Non significa abbandonare gli altri trentasette comuni. Significa riconoscere che, quando il tempo e le persone sono limitati, bisogna partire dai luoghi in cui la scadenza politica è più vicina e dove una presenza organizzata può produrre risultati visibili.")
    para(doc, "Il dato di circa quattrocentotrentacinquemila cittadini coinvolti non indica il numero dei possibili elettori di Futuro Nazionale. È una stima della popolazione complessivamente interessata dalle amministrative nei diciotto comuni. Serve a far capire che non si sta parlando di un’operazione marginale: le città e i paesi coinvolti rappresentano una parte molto rilevante della provincia.")
    para(doc, "I nove mesi indicati in copertina sono l’orizzonte di lavoro disponibile al momento della preparazione del progetto. Non sono nove mesi di propaganda continua. Sono il tempo necessario per ascoltare, organizzare, formare, scegliere le priorità, decidere dove e come presentarsi e, solo alla fine, affrontare la campagna elettorale vera e propria.")
    note(doc, "Il messaggio politico", "La Brianza non ha bisogno di una sigla che compare poche settimane prima del voto. Ha bisogno di una struttura che inizi ora, lavori sul territorio e sia in grado di dimostrare che cosa ha costruito.")


def page_luca(doc):
    section_title(doc, "Perché la proposta ruota intorno a Luca", "Profilo e credibilità")
    lead(doc, "Il dossier non presenta Luca come un personaggio inventato per l’occasione. Lo presenta come una persona che possiede quattro risorse utili alla fase in cui si trova il movimento: esperienza pubblica, competenza, tempo disponibile e capacità di lavorare con una squadra.")
    para(doc, "Quando si ricorda che ha ricoperto ruoli di consigliere provinciale e di assessore comunale, non si vuole costruire un curriculum celebrativo. Si vuole dire che conosce i meccanismi reali degli enti locali: sa che una proposta comunale deve confrontarsi con bilanci, competenze, regolamenti, tempi amministrativi e rapporti tra maggioranza e opposizione. Questa conoscenza riduce il rischio di produrre slogan impossibili da trasformare in azioni.")
    para(doc, "Il fatto che sia architetto aggiunge una capacità di lettura del territorio. Un architetto è abituato a osservare spazi, infrastrutture, trasformazioni urbane e vincoli. Il fatto che abbia insegnato per molti anni negli istituti superiori aggiunge un’altra competenza: spiegare questioni complesse, ascoltare persone diverse, mantenere un metodo e accompagnare la crescita di altri.")
    para(doc, "La pensione viene presentata come un vantaggio operativo, non come un elemento anagrafico. Molte organizzazioni territoriali funzionano soltanto la sera perché chi le guida ha un lavoro a tempo pieno. Luca, invece, può incontrare amministratori, associazioni, commercianti e cittadini anche durante il giorno. Può seguire telefonate, riunioni e problemi con continuità. In una fase di costruzione, la disponibilità di tempo è una risorsa politica concreta.")
    para(doc, "Infine c’è il tema della squadra. Luca non propone di accentrare ogni compito. Propone di assegnare deleghe vere, con obiettivi e verifiche. Questo punto serve anche a rassicurare chi teme di essere escluso. Entrare nella struttura non dovrebbe significare ricevere un titolo ornamentale, ma assumersi una parte di lavoro riconoscibile.")
    note(doc, "In parole semplici", "La candidatura di Luca vuole essere credibile non perché promette di sapere tutto, ma perché combina esperienza, metodo, tempo e disponibilità a condividere responsabilità.", fill="FFF7D9", accent=GOLD)


def page_mandate(doc):
    section_title(doc, "Che cosa significa diventare Coordinatore Provinciale", "Mandato e Statuto")
    lead(doc, "La prima precisazione del dossier è anche una delle più importanti: il nome corretto della funzione è Coordinatore Provinciale. Usare il titolo previsto dallo Statuto dimostra rispetto delle regole e impedisce di presentare la proposta come una struttura personale costruita fuori dal movimento.")
    para(doc, "Nella fase transitoria, la scelta del coordinatore provinciale avviene secondo le procedure interne previste dal movimento. In seguito, quando la struttura sarà pienamente costituita, il ruolo sarà legato al congresso provinciale e avrà la durata indicata dallo Statuto. Il dossier non tenta di sostituire queste procedure. Chiede che Luca venga valutato come persona capace di svolgere quel mandato.")
    para(doc, "Il coordinatore non è soltanto chi rappresenta il movimento nelle fotografie o nelle riunioni. È la persona che tiene insieme indirizzo politico, organizzazione, rapporti con il livello regionale e nazionale, gestione delle differenze interne e presenza nei comuni. Deve sapere quando decidere direttamente, quando delegare e quando chiedere una decisione agli organi superiori.")
    para(doc, "Il riferimento ai valori V.I.T.A.L.E. serve a collegare il metodo locale all’identità nazionale del movimento. Virtù, Identità, Tradizioni, Amore, Libertà ed Eccellenza non vengono presentati come parole decorative. Nel dossier ciascun valore viene tradotto in un comportamento: responsabilità personale, coerenza, rispetto delle radici, attenzione alla comunità, chiarezza nelle decisioni ed esigenza di qualità.")
    para(doc, "Questo collegamento è essenziale perché una struttura provinciale non può diventare un partito nel partito. Può adattare il lavoro alla Brianza, ma deve restare riconoscibile dentro lo Statuto, il manifesto e il programma di Futuro Nazionale.")
    note(doc, "Il mandato vero", "Luca non chiede libertà assoluta. Chiede l’autorità necessaria per costruire, dentro regole già esistenti, una struttura provinciale che possa essere controllata e corretta.")


def page_organization(doc):
    section_title(doc, "Come sarebbe organizzata la squadra", "Architettura provinciale")
    lead(doc, "L’organigramma proposto distingue la direzione politica dalle funzioni operative. È un modo per evitare due errori opposti: che tutto dipenda da una sola persona oppure che molti abbiano un titolo ma nessuno sappia chi deve rispondere di un risultato.")
    para(doc, "Al vertice c’è il Coordinatore Provinciale. A lui spettano l’indirizzo politico, la rappresentanza e la responsabilità complessiva. Accanto a lui opera un Esecutivo Provinciale, cioè un gruppo ristretto che traduce l’indirizzo in decisioni periodiche e controlla lo stato del lavoro. Il numero e la composizione definitiva devono restare coerenti con lo Statuto.")
    para(doc, "La funzione organizzativa segue eventi, sedi, calendario, logistica e copertura dei comuni. La funzione tesseramento segue iscrizioni, rinnovi, qualità e integrità dei dati. La direzione disciplinare interviene quando ci sono regole da applicare o comportamenti incompatibili con il movimento. I laboratori tematici preparano proposte politiche; il settore giovani collega scuole, università e nuove generazioni; la segreteria garantisce verbali, agenda, documenti e continuità amministrativa.")
    para(doc, "Le deleghe operative servono a rendere possibile il lavoro quotidiano. Nel dossier compaiono funzioni come rapporti con enti locali, amministrative 2027, programma comunale, comunicazione, eventi, dati e privacy. Non compaiono nomi perché questa fase non deve sembrare una spartizione preventiva di incarichi. Prima si definiscono i compiti, poi si individuano le persone più adatte.")
    para(doc, "La frase ‘una squadra larga, con responsabilità nette’ riassume l’equilibrio cercato. Larga significa che devono esserci spazi per competenze e disponibilità diverse. Responsabilità nette significa che, quando un’attività non viene fatta, si deve sapere chi aveva il compito, quale ostacolo ha incontrato e come correggere il problema.")
    note(doc, "Per chi teme di essere escluso", "La proposta non chiude la porta a nessuno che voglia lavorare seriamente. Chiede però che ogni ruolo sia accompagnato da un obiettivo, una scadenza e una verifica.", fill="FFF7D9", accent=GOLD)


def page_delegation(doc):
    section_title(doc, "Delegare senza creare piccoli feudi", "Inclusione e responsabilità")
    lead(doc, "In politica la parola ‘delega’ può essere interpretata come distribuzione di potere. Qui viene usata in un senso più concreto: distribuzione di lavoro, informazioni e responsabilità, mantenendo una direzione provinciale comune.")
    para(doc, "Un referente d’area non diventa il proprietario dei comuni che segue. Deve raccogliere informazioni, favorire incontri, sostenere i referenti locali e riferire alla struttura provinciale. Allo stesso modo, chi segue la comunicazione non decide autonomamente la linea politica; trasforma in contenuti una linea approvata e segnala quando mancano informazioni o autorizzazioni.")
    para(doc, "La proposta prevede che le nomine siano coerenti con lo Statuto, che gli obiettivi siano misurabili e che le persone vengano accompagnate da formazione e verifiche. Questo evita che un incarico diventi permanente soltanto perché è stato assegnato una volta. Se una funzione non produce risultati, il problema deve essere affrontato apertamente e la delega può essere corretta.")
    para(doc, "L’inclusione non viene quindi misurata dal numero di titoli concessi, ma dalla possibilità reale di partecipare. Una persona può contribuire con relazioni istituzionali, organizzazione di eventi, conoscenza di un comune, competenze professionali, produzione di documenti, comunicazione o attività di volontariato. Il compito della guida è trasformare queste disponibilità in un sistema ordinato.")
    para(doc, "Il coordinatore conserva la responsabilità finale. Non può attribuire un compito e poi disinteressarsene. Deve ricevere report, rimuovere ostacoli, intervenire nei conflitti e riferire al livello regionale o nazionale quando una decisione supera la sua competenza.")
    note(doc, "La regola pratica", "Delegare significa dire con chiarezza: questo è il risultato atteso, questa è la persona responsabile, questa è la data di controllo. Non significa soltanto assegnare un titolo.")


def page_areas(doc):
    section_title(doc, "Perché dividere la provincia in cinque aree", "Presenza territoriale")
    lead(doc, "Cinquantacinque comuni sono troppi per essere seguiti direttamente ogni giorno da una sola segreteria. Le cinque macroaree sono quindi uno strumento di lavoro: avvicinano il coordinamento ai territori senza spezzare la provincia in cinque organizzazioni indipendenti.")
    para(doc, "La prima area ruota intorno a Monza e alla cintura vicina. La seconda comprende l’asse del Seveso e comuni come Cesano Maderno, Meda, Seveso e Varedo. La terza guarda alla Brianza centrale, la quarta alla Valle del Lambro e la quinta al Vimercatese. La composizione esatta può essere corretta dopo il censimento dei referenti e delle realtà già attive.")
    para(doc, "Ogni area dovrebbe produrre quattro cose semplici. La prima è un’agenda minima, cioè almeno un’iniziativa o un coordinamento periodico. La seconda è la copertura: bisogna sapere quali comuni hanno un referente vero e quali invece sono ancora scoperti. La terza è l’ascolto: raccogliere temi locali documentati, non soltanto impressioni. La quarta è il rendiconto: una breve nota mensile che dica che cosa è stato fatto e che cosa non sta funzionando.")
    para(doc, "La macroarea non sostituisce il comune. Serve a sostenere i comuni più piccoli, che spesso non hanno abbastanza persone per mantenere una struttura autonoma, e a evitare duplicazioni. Per esempio, più comuni vicini possono organizzare insieme una serata sulla mobilità, sulla sanità o sulla sicurezza, mantenendo poi una lettura specifica dei problemi locali.")
    para(doc, "La presenza territoriale deve essere regolare. Un partito che visita un paese soltanto prima delle elezioni appare opportunista. Un coordinamento che torna, riferisce e mantiene gli impegni costruisce invece reputazione, anche quando non dispone ancora di una lista o di un candidato.")
    note(doc, "Una sola provincia", "Le cinque aree distribuiscono il lavoro, ma messaggio, regole, dati e responsabilità restano provinciali. Nessuna area deve diventare una corrente o una struttura personale.")


def page_100_days(doc):
    section_title(doc, "Che cosa deve succedere nei primi cento giorni", "Cronoprogramma")
    lead(doc, "I cento giorni non sono una formula pubblicitaria. Sono un periodo abbastanza breve da creare urgenza e abbastanza lungo da consentire un censimento serio, la definizione dei ruoli e l’avvio delle prime attività.")
    para(doc, "Nei primi quindici giorni la priorità è ascoltare. Luca dovrebbe incontrare comitati, amministratori, referenti, persone disponibili e realtà già presenti. Non deve partire imponendo un organigramma definitivo, perché prima bisogna capire chi esiste, che cosa sa fare, quali rapporti sono attivi e dove ci sono tensioni o vuoti.")
    para(doc, "Tra il quindicesimo e il trentesimo giorno si costruisce la mappa: comuni coperti e scoperti, referenti, competenze, organizzazioni locali e stato dei diciotto comuni al voto. Questa mappa è la base delle decisioni successive. Senza di essa si rischia di distribuire incarichi sulla base delle sole relazioni personali.")
    para(doc, "Tra il trentesimo e il sessantesimo giorno si organizza la macchina. Si formalizzano le funzioni operative, si avviano le cinque macroaree, si stabiliscono calendario, regole di lavoro, comunicazione e flussi informativi. In questa fase devono diventare chiare anche le modalità con cui vengono raccolti e protetti i dati degli iscritti e dei contatti.")
    para(doc, "Tra il sessantesimo e il novantesimo giorno si preparano i dossier dei diciotto comuni. Ogni dossier deve contenere situazione politica, problemi principali, persone disponibili, possibili formule elettorali, relazioni e rischi. Dal novantesimo al centesimo giorno si esegue una prima verifica: che cosa è stato realmente costruito, dove esistono ritardi e quali decisioni devono salire al livello regionale o nazionale.")
    para(doc, "Il cruscotto mensile controlla copertura, radicamento, attività, preparazione, disciplina e qualità. Non serve a punire chi lavora, ma a rendere visibili i problemi prima che diventino emergenze.")
    note(doc, "La promessa verificabile", "Dopo cento giorni non si pretende che tutti i cinquantacinque comuni abbiano una sezione completa. Si pretende però una mappa affidabile, una struttura leggibile e un piano concreto per colmare i vuoti.")


def page_labs(doc):
    section_title(doc, "Gli otto laboratori politici", "Dall’organizzazione alle proposte")
    lead(doc, "Una struttura territoriale non può limitarsi a organizzare cene, riunioni e banchetti. Deve trasformare l’ascolto in proposte. Per questo il dossier prevede otto laboratori tematici, ciascuno collegato a problemi riconoscibili della Brianza.")
    para(doc, "Il laboratorio su sicurezza e legalità affronta presìdi, degrado, commercio, integrazione tra servizi e polizia locale. Quello su impresa e lavoro riguarda piccole e medie imprese, artigianato, burocrazia, formazione e attrattività. Famiglia e comunità comprende scuola, servizi, natalità, volontariato e fragilità. Scuola e giovani guarda a merito, orientamento, istituti tecnici superiori, sport e cittadinanza.")
    para(doc, "Il laboratorio su territorio e ambiente si occupa di urbanistica, consumo di suolo, bonifiche, acqua e parchi. Mobilità e infrastrutture affronta ferrovie, strade, trasporto locale e collegamenti tra est e ovest. Identità e cultura raccoglie storia locale, tradizioni, patrimonio, biblioteche ed eventi. Salute e sociale riguarda medicina territoriale, disabilità, tempi di accesso e integrazione dei servizi.")
    para(doc, "Il metodo è importante quanto i temi. Una proposta deve partire da un problema osservabile e dalla competenza dell’ente che potrebbe intervenire. Se il problema è comunale, non bisogna fingere che possa risolverlo la Regione; se è regionale o nazionale, bisogna individuare l’interlocutore corretto. Poi si consultano amministratori, professionisti, associazioni e cittadini, si separano il principio politico dalla soluzione tecnica e si pubblica soltanto ciò che è stato verificato e approvato.")
    para(doc, "L’obiettivo, entro la primavera del 2027, è avere per ogni comune prioritario almeno una scheda sintetica che contenga problema, dati disponibili, competenza istituzionale, soluzione proposta, vincoli e costo indicativo quando conoscibile. In questo modo il programma non nasce copiando slogan nazionali, ma leggendo la Brianza dentro i valori del movimento.")
    note(doc, "Il vantaggio", "Un laboratorio ben gestito diventa una piccola fabbrica di proposte e forma nuove persone. Chi possiede competenze può contribuire anche senza ricoprire subito un incarico politico.")


def page_governance(doc):
    section_title(doc, "Le regole che proteggono la squadra", "Governo, conflitti e reputazione")
    lead(doc, "La parte sulle regole può sembrare meno attraente di quella elettorale, ma è fondamentale. Le organizzazioni giovani spesso crescono rapidamente e poi si bloccano per conflitti personali, aspettative non chiarite, gestione opaca delle risorse o incarichi interpretati come proprietà.")
    para(doc, "Le decisioni operative e reversibili possono essere delegate. Un calendario, una riunione o un formato di report possono essere corretti senza danni rilevanti. Le decisioni politiche difficili da annullare, invece, devono restare agli organi competenti: accordi elettorali, candidature, spese importanti, incarichi, posizioni pubbliche e comunicazioni che possono produrre conseguenze reputazionali.")
    para(doc, "Per ogni attività deve esistere una prova. Non basta dire che un comune è coperto: bisogna indicare un referente effettivo, un contatto e un’attività. Non basta dire che una riunione è andata bene: bisogna registrare presenze, decisioni e prossime azioni. Questa disciplina aiuta anche il coordinatore, perché gli consente di riferire con precisione al livello regionale e nazionale.")
    para(doc, "Il dossier stabilisce inoltre una linea etica netta: nessuna promessa di compensazioni, posti in società partecipate, nomine o vantaggi personali. Gli incarichi devono servire al movimento e rispettare competenza, trasparenza e assenza di conflitti di interesse. Questo punto protegge Luca e l’intera struttura da aspettative improprie.")
    para(doc, "Quando nasce un conflitto, non lo si deve coprire con il silenzio né trasformarlo immediatamente in una battaglia pubblica. Si chiariscono fatti, regole e competenze, si cerca una soluzione interna e, se necessario, si porta la questione all’organo superiore. La reputazione provinciale dipenderà molto dalla capacità di gestire le differenze senza vendette e senza disordine.")
    note(doc, "La linea non negoziabile", "Nessun ruolo organizzativo deve essere presentato come anticipo di incarichi pubblici, candidature o benefici futuri.", fill=RED_LIGHT, accent=RED)


def page_numbers(doc):
    section_title(doc, "Come leggere i numeri elettorali", "Spiegazione per non esperti")
    lead(doc, "Le pagine colorate non dicono quanti voti prenderà Futuro Nazionale. Descrivono il terreno sul quale il movimento dovrebbe lavorare. Per usarle correttamente bisogna distinguere tre informazioni diverse: amministrazione comunale, voto europeo e affluenza.")
    para(doc, "Il colore blu indica che l’ultima amministrazione viene ricondotta al centrodestra; il rosso indica il centrosinistra; il verde acqua segnala una situazione civica; il grigio una gestione commissariale. La percentuale grande è quella ottenuta dal candidato sindaco vincente nell’ultima elezione comunale disponibile. Non rappresenta la somma dei partiti nazionali: alle comunali contano molto il candidato, le liste civiche, le alleanze e le questioni locali.")
    para(doc, "Le tre cifre più piccole vengono dalle elezioni europee del 2024. L’area CDX è la somma dei voti ottenuti da Fratelli d’Italia, Lega e Forza Italia. L’area CSX è la somma di Partito Democratico e Alleanza Verdi e Sinistra. La terza cifra è l’affluenza, cioè la percentuale degli aventi diritto che ha votato. Mancano quindi altri partiti e liste: per questo CDX e CSX non sommano al cento per cento.")
    para(doc, "Nei diciotto comuni considerati, l’aggregato europeo assegna circa il 50,8 per cento ai tre partiti dell’area di centrodestra e circa il 31 per cento a PD e AVS. Tra le singole liste, Fratelli d’Italia è intorno al 30,3 per cento, il PD al 23,5, Forza Italia al 10,9, la Lega al 9,7, AVS al 7,5 e il Movimento 5 Stelle al 6,1. Questi dati mostrano un bacino complessivamente favorevole al centrodestra, ma non dicono come quel bacino si distribuirà nel 2027.")
    para(doc, "Futuro Nazionale non era presente sulla scheda europea del 2024. Non è quindi corretto attribuirgli automaticamente una parte dei voti di Fratelli d’Italia, Lega o Forza Italia. I numeri servono soltanto a capire dove esiste un elettorato di area più forte, dove il centrosinistra è competitivo e dove la partecipazione è più bassa.")
    note(doc, "La cautela fondamentale", "Europee e comunali sono elezioni diverse. Un comune amministrato dal centrosinistra può avere un voto europeo prevalentemente di centrodestra, e viceversa. Questa differenza è un tema da studiare, non una vittoria già disponibile.", fill="FFF7D9", accent=GOLD)


TOWNS = {
    "Arcore": ("centrodestra", "50,86%", "48,9%", "33,2%", "53,7%", "Il risultato comunale è stato vicino alla maggioranza assoluta. Il voto europeo mostra un’area di centrodestra importante, ma non dominante: serve una presenza capace di distinguersi senza dare per acquisito il consenso."),
    "Biassono": ("centrodestra", "43,32%", "56,5%", "26,4%", "56,3%", "La percentuale del sindaco vincente sotto il cinquanta per cento suggerisce una competizione con più liste. Alle Europee il centrodestra è forte. La priorità è capire gli equilibri locali e le relazioni già consolidate."),
    "Bovisio-Masciago": ("commissariato", "50,60% nel 2024", "52,9%", "29,8%", "59,6%", "Il comune richiede cautela perché la situazione amministrativa è cambiata dopo lo scioglimento del consiglio. I numeri precedenti restano una fotografia storica, ma l’analisi politica deve partire dal nuovo contesto."),
    "Briosco": ("centrodestra", "54,46%", "62,0%", "24,3%", "53,6%", "È uno dei bacini europei più forti per il centrodestra. Questo non rende facile l’ingresso di un nuovo movimento: significa anche confrontarsi con partiti e amministratori già radicati."),
    "Carnate": ("centrodestra", "39,38%", "41,9%", "38,5%", "54,6%", "Il risultato comunale mostra una frammentazione significativa. Alle Europee le due aree sono relativamente vicine. Qui contano molto candidato, liste civiche e capacità di costruire relazioni locali."),
    "Cesano Maderno": ("centrosinistra", "53,94%", "55,7%", "27,0%", "48,1%", "L’amministrazione è di centrosinistra, mentre alle Europee il centrodestra risulta ampiamente avanti. La differenza dimostra quanto la qualità dell’offerta comunale possa cambiare il comportamento degli elettori."),
    "Lentate sul Seveso": ("centrodestra", "67,03%", "57,8%", "27,2%", "49,2%", "L’ultima amministrativa è stata nettamente favorevole al centrodestra. Una nuova presenza deve portare competenze e temi riconoscibili, non limitarsi a dichiarare appartenenza alla stessa area."),
    "Lesmo": ("centrosinistra", "33,38%", "56,3%", "27,5%", "59,4%", "La percentuale del vincitore indica una competizione frammentata. Il voto europeo è invece chiaramente orientato al centrodestra. Prima di parlare di opportunità occorre comprendere liste civiche e rapporti personali."),
    "Limbiate": ("centrodestra", "71,27%", "53,6%", "28,5%", "43,6%", "La vittoria comunale è stata molto ampia, ma l’affluenza europea è la più bassa del gruppo. Il problema strategico non è soltanto convincere, ma ricostruire partecipazione e presenza."),
    "Lissone": ("centrodestra", "50,43%", "54,8%", "27,9%", "49,9%", "Il risultato comunale è stato molto vicino alla soglia del cinquanta per cento. È una città grande e strategica, dove servono struttura, continuità e una proposta locale credibile."),
    "Meda": ("centrodestra", "68,50%", "59,1%", "24,7%", "51,2%", "Il centrodestra appare solido sia sul piano comunale sia su quello europeo. La sfida consiste nel dimostrare quale valore aggiunto possa offrire Futuro Nazionale in un’area già politicamente presidiata."),
    "Monza": ("centrosinistra", "51,21%", "44,4%", "35,0%", "52,7%", "È il capoluogo e ha un peso politico particolare. Le due aree europee sono più vicine rispetto a molti altri comuni. Per essere credibili servono temi cittadini, persone competenti e una presenza che non appaia soltanto provinciale."),
    "Seveso": ("centrodestra", "53,52%", "57,7%", "26,0%", "46,1%", "Il centrodestra dispone di un bacino forte, mentre la partecipazione europea è contenuta. Il radicamento dovrebbe unire questioni identitarie, ambiente, sicurezza e capacità di riportare persone alla vita pubblica."),
    "Sulbiate": ("civica", "53,82%", "51,6%", "31,1%", "50,3%", "La guida civica invita a non leggere il paese soltanto con categorie nazionali. Il voto europeo è favorevole al centrodestra, ma la formula comunale dipenderà molto da persone e progetti locali."),
    "Varedo": ("centrodestra", "68,83%", "52,5%", "29,1%", "52,0%", "La maggioranza comunale è stata larga. Anche qui la semplice appartenenza di area non basta: bisogna capire se esistono spazi politici, competenze e relazioni che giustifichino un investimento autonomo."),
    "Vedano al Lambro": ("centrodestra", "55,88%", "51,2%", "31,2%", "58,5%", "Il comune presenta un centrodestra maggioritario ma non schiacciante e una partecipazione europea relativamente alta. È adatto a un lavoro di ascolto e presenza ben calibrato."),
    "Verano Brianza": ("centrosinistra", "49,56%", "58,2%", "26,4%", "53,7%", "La guida comunale è di centrosinistra, mentre alle Europee l’area di centrodestra è molto forte. La distanza segnala l’importanza decisiva del candidato e delle dinamiche civiche."),
    "Vimercate": ("centrosinistra", "60,87%", "39,9%", "41,2%", "55,1%", "È il comune del gruppo in cui, alle Europee, PD e AVS insieme superano di poco i tre partiti del centrodestra. Per crescere occorre una proposta seria su servizi, mobilità, impresa, ambiente e identità locale."),
}


def town_block(doc, name):
    orientation, municipal, cdx, csx, turnout, reading = TOWNS[name]
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    B.set_run(p.add_run(name), size=11.2, color=NAVY, bold=True)
    para(doc, f"Ultima amministrazione: {orientation}; risultato del vincitore {municipal}. Europee 2024: area CDX {cdx}, area CSX {csx}, affluenza {turnout}. {reading}", size=9.55, after=6, line=1.16, first=False)


def page_towns(doc, names, number):
    section_title(doc, f"I comuni spiegati uno per uno · {number}/3", "Lettura territoriale")
    lead(doc, "Ogni scheda mette insieme dati diversi per suggerire domande di lavoro. Le brevi interpretazioni seguenti non indicano dove il movimento prenderà più voti: indicano che cosa bisogna verificare prima di decidere priorità, alleanze o candidature.")
    for name in names:
        town_block(doc, name)


def page_strategy(doc):
    section_title(doc, "Dai numeri alle priorità", "Lettura strategica")
    lead(doc, "Il dossier raggruppa i comuni non per assegnare vincitori in anticipo, ma per capire quale tipo di lavoro organizzativo è più urgente. Lo stesso metodo non può essere applicato indistintamente a Monza, a un comune civico o a un paese con una maggioranza di centrodestra molto consolidata.")
    para(doc, "Nei bacini dove il centrodestra europeo è particolarmente forte, come Briosco, Meda, Verano, Lentate e Seveso, esiste un pubblico potenzialmente sensibile ai temi del movimento. Tuttavia esiste anche una concorrenza politica già strutturata. Futuro Nazionale deve spiegare che cosa aggiunge, evitando di comportarsi come se tutti i voti d’area fossero trasferibili.")
    para(doc, "Nei comuni amministrati dal centrosinistra ma con un voto europeo favorevole al centrodestra, come Cesano Maderno e Verano Brianza, la domanda principale è perché l’area non sia riuscita a trasformare il proprio bacino nazionale in una proposta comunale vincente. Le possibili risposte riguardano candidati, divisioni, liste civiche, credibilità e temi locali.")
    para(doc, "Monza e Vimercate richiedono un approccio specifico. Sono centri di peso, con sistemi politici e sociali più complessi. Non basta organizzare un banchetto occasionale: servono competenze, relazioni, contenuti cittadini e continuità. Anche Lissone, Limbiate, Meda, Cesano e Seveso hanno dimensioni che richiedono una squadra locale reale.")
    para(doc, "I comuni più piccoli non devono essere considerati secondari. Possono offrire radicamento, amministratori, volontari e capacità di sperimentare metodi che poi vengono replicati altrove. In questi luoghi le relazioni personali e le liste civiche hanno spesso un peso superiore alle etichette nazionali.")
    para(doc, "Le priorità finali devono essere decise dopo aver aggiunto ai numeri elettorali cinque informazioni che oggi mancano o sono incomplete: iscritti effettivi, referenti affidabili, candidato possibile, accesso alle coalizioni e problemi locali capaci di mobilitare partecipazione.")
    note(doc, "La decisione prudente", "Prima si misura la forza organizzativa reale, poi si decide se presentare una lista, entrare in una coalizione, sostenere un progetto civico compatibile oppure limitarsi a costruire presenza.", fill="FFF7D9", accent=GOLD)


def page_ask(doc):
    section_title(doc, "Che cosa Luca chiede realmente", "La proposta al tavolo")
    lead(doc, "Il dossier non chiede soltanto una nomina. Chiede un mandato operativo accompagnato da tempi, strumenti e possibilità di verifica. È questa la differenza tra occupare una posizione e assumersi la responsabilità di costruire.")
    para(doc, "Il primo punto è il riconoscimento del ruolo di Coordinatore Provinciale secondo le procedure interne. Il secondo è un periodo di novanta-cento giorni per completare censimento, organizzazione e dossier dei comuni. Il terzo è la possibilità di proporre deleghe e referenti d’area senza anticipare nomi e rispettando gli organi statutari. Il quarto è l’accesso ai dati e ai contatti strettamente necessari, sempre nel rispetto della privacy.")
    para(doc, "Le prime tre mosse dopo la nomina hanno un significato preciso. Entro quarantotto ore Luca dovrebbe inviare un messaggio unitario alle realtà esistenti, facendo capire che non arriva per cancellare il lavoro di altri. Entro sette giorni dovrebbe convocare una prima riunione operativa per costruire la mappa delle disponibilità. Entro quindici giorni dovrebbe produrre una nota al livello regionale o nazionale con copertura, rischi e decisioni necessarie.")
    para(doc, "La frase centrale del discorso è che la Brianza non ha bisogno di un coordinatore che occupi uno spazio, ma di una guida che costruisca una squadra. L’esperienza amministrativa e provinciale, la competenza tecnica e la disponibilità di tempo diventano credibili soltanto se vengono legate a risultati: cinque aree operative, diciotto dossier, responsabilità distribuite e rendicontazione.")
    para(doc, "Luca non deve presentare tutto questo come un sistema già imposto. Deve presentarlo come una base di lavoro aperta ai contributi. Il mandato gli darebbe la responsabilità di guidare il processo, non il diritto di decidere in solitudine ogni dettaglio.")
    note(doc, "La formula più chiara", "Non vi chiedo un titolo da esibire. Vi propongo un mandato da verificare: entro cento giorni saprete che cosa è stato costruito, dove siamo presenti e quali decisioni servono per il 2027.")


def page_table_use(doc):
    section_title(doc, "Come portare il dossier alla riunione", "Uso pratico")
    lead(doc, "Il documento è forte se viene usato come prova di preparazione. Diventa debole se Luca tenta di leggerlo integralmente o se lo presenta come un piano già chiuso, scritto senza bisogno degli altri.")
    para(doc, "Luca dovrebbe parlare prima con parole proprie. Può spiegare in circa due minuti che vuole costruire una squadra provinciale, che ha studiato le scadenze del 2027 e che propone un metodo verificabile. Solo dopo dovrebbe mostrare la copertina, aprire la sintesi, far vedere il cronoprogramma e indicare una pagina con i comuni.")
    para(doc, "La parte numerica serve a dimostrare che la proposta non è generica. Non è necessario commentare tutte le percentuali. Basta spiegare che sono state incrociate le ultime amministrative con le Europee 2024 per capire dove approfondire l’analisi. Se qualcuno chiede dettagli, Luca può rispondere usando questa guida o la pagina personale delle fonti.")
    para(doc, "Il tono deve essere inclusivo. Espressioni come ‘questa è la struttura che farò’ possono allarmare chi teme di perdere spazio. È meglio dire ‘questa è una base sulla quale coinvolgere le persone disponibili, definendo prima i compiti e poi i nomi’. In questo modo la preparazione appare come una garanzia, non come un tentativo di occupazione.")
    para(doc, "Dopo l’intervento, senza insistere, Luca può consegnare una copia delle prime quattordici pagine a chi deve decidere. Il promemoria personale e le spiegazioni operative restano a lui. Il dossier non deve essere distribuito indiscriminatamente prima che siano chiariti mandato e interlocutori.")
    para(doc, "La forza della presentazione sta nell’equilibrio: mostrare ambizione senza arroganza, competenza senza tecnicismo, disponibilità alla squadra senza rinunciare alla direzione. Se questo equilibrio viene mantenuto, i fogli diventano la prova che Luca è già entrato mentalmente nel lavoro del coordinatore.")
    note(doc, "La frase di consegna", "Ho provato a mettere tutto per iscritto, con tempi, dati e responsabilità. È una proposta aperta, ma permette di capire da dove partire già il giorno dopo la decisione.", fill="FFF7D9", accent=GOLD)


def page_glossary(doc):
    section_title(doc, "Parole e cautele da ricordare", "Glossario finale")
    lead(doc, "Questa pagina raccoglie i termini che possono creare equivoci e le precauzioni necessarie per usare il documento in modo corretto.")
    entries = [
        ("CDX", "Nel dossier è la somma analitica di Fratelli d’Italia, Lega e Forza Italia alle Europee 2024; non è una coalizione comunale già definita."),
        ("CSX", "Nel dossier è la somma analitica di Partito Democratico e Alleanza Verdi e Sinistra; non comprende automaticamente tutte le liste di centrosinistra."),
        ("Affluenza", "Percentuale degli aventi diritto che ha votato. Una bassa affluenza può indicare disinteresse, difficoltà di mobilitazione o caratteristiche specifiche dell’elezione."),
        ("Commissariamento", "Gestione temporanea del comune da parte di un commissario dopo lo scioglimento degli organi eletti. Rende meno utile leggere il futuro soltanto attraverso l’ultima maggioranza."),
        ("Macroarea", "Raggruppamento operativo di comuni vicini. Non è una struttura politica autonoma e non sostituisce la provincia."),
        ("Dossier comunale", "Scheda di lavoro con situazione politica, temi, persone, relazioni, rischi e decisioni. Non è necessariamente un programma pubblico."),
        ("Indicatore", "Numero o informazione che permette di capire se il lavoro procede. Non sostituisce il giudizio politico, ma lo rende più informato."),
        ("Porta a una via", "Decisione difficile da annullare, come candidatura, alleanza, spesa o comunicazione pubblica delicata. Richiede approvazione e prove maggiori."),
    ]
    table = doc.add_table(rows=len(entries), cols=2)
    B.set_table_fixed(table, [3.7, 14.1])
    for row, (label, body) in zip(table.rows, entries):
        B.shade(row.cells[0], NAVY)
        B.shade(row.cells[1], LIGHT)
        B.borders(row.cells[0], WHITE, 8)
        B.borders(row.cells[1], WHITE, 8)
        B.cell_text(row.cells[0], label, size=8.4, color=WHITE, bold=True)
        B.cell_text(row.cells[1], body, size=8.3)
    para(doc, "La data delle amministrative 2027 non era ancora ufficialmente fissata al momento della redazione. Maggio è una finestra di pianificazione da aggiornare dopo il decreto di convocazione. Anche composizione delle liste, alleanze, candidature, dati associativi e struttura riconosciuta devono essere aggiornati dopo la riunione e le decisioni degli organi competenti.", size=8.8, color=MUTED, italic=True, first=False)
    note(doc, "Conclusione", "Il dossier è utile perché non finge di sapere già tutto. Mostra che Luca conosce il lavoro da fare, propone un metodo e indica con chiarezza quali informazioni e decisioni mancano ancora.")


def configure(doc):
    B.configure_document(doc)
    sec = doc.sections[0]
    header = sec.header
    header.paragraphs[0].text = ""
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    B.set_run(hp.add_run("GUIDA DISCORSIVA · PROPOSTA LUCA VIVIANI"), size=7.2, color=MUTED, bold=True)
    footer = sec.footer
    footer.paragraphs[0].text = ""
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    B.set_run(fp.add_run("DOCUMENTO DI ACCOMPAGNAMENTO · USO INTERNO · 29.08.2026"), size=7.0, color=MUTED)
    # Narrative-proposal rhythm, with A4 geometry inherited from the original dossier.
    normal = doc.styles["Normal"]
    normal.paragraph_format.line_spacing = 1.22
    normal.paragraph_format.space_after = Pt(7)


def audit(doc):
    assert len(doc.sections) == 1
    assert round(doc.sections[0].page_width.cm, 1) == 21.0
    assert round(doc.sections[0].page_height.cm, 1) == 29.7
    assert len(doc.inline_shapes) == 1
    for table in doc.tables:
        assert table.autofit is False
        for row in table.rows:
            assert row.height is None
    text = "\n".join(p.text for p in doc.paragraphs).lower()
    for phrase in ["cinquantacinque comuni", "Coordinatore Provinciale", "Europee del 2024", "Luca non deve"]:
        assert phrase.lower() in text


def build():
    logo = B.crop_logo()
    doc = Document()
    configure(doc)
    pages = [
        lambda: page_cover(doc, logo),
        lambda: page_premessa(doc),
        lambda: page_big_picture(doc),
        lambda: page_luca(doc),
        lambda: page_mandate(doc),
        lambda: page_organization(doc),
        lambda: page_delegation(doc),
        lambda: page_areas(doc),
        lambda: page_100_days(doc),
        lambda: page_labs(doc),
        lambda: page_governance(doc),
        lambda: page_numbers(doc),
        lambda: page_towns(doc, ["Arcore", "Biassono", "Bovisio-Masciago", "Briosco", "Carnate", "Cesano Maderno"], 1),
        lambda: page_towns(doc, ["Lentate sul Seveso", "Lesmo", "Limbiate", "Lissone", "Meda", "Monza"], 2),
        lambda: page_towns(doc, ["Seveso", "Sulbiate", "Varedo", "Vedano al Lambro", "Verano Brianza", "Vimercate"], 3),
        lambda: page_strategy(doc),
        lambda: page_ask(doc),
        lambda: page_table_use(doc),
        lambda: page_glossary(doc),
    ]
    for i, fn in enumerate(pages):
        if i:
            B.add_page_break(doc)
        fn()
    audit(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
